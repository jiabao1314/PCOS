"""Build the user-facing PCOS comment-by-comment execution plan.
Scientific outputs are read-only inputs; this script performs no scientific analysis.
"""
import json
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

B=Path(__file__).resolve().parent
DATA=json.loads((B/'20261030_review_plan_content.json').read_text())
CROSS=json.loads((B/'20261030_v5_comment_crosswalk.json').read_text())
OUT=B/'20261030_PCOS_17条批注修改与Codex执行方案.docx'
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.72);s.bottom_margin=Inches(.68)
s.left_margin=Inches(.85);s.right_margin=Inches(.85)
s.header_distance=Inches(.3);s.footer_distance=Inches(.3)

def fontstyle(style,size,bold=False):
    f=style.font;f.name='Liberation Sans';f.size=Pt(size);f.bold=bold;f.italic=False;f.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr(); rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for k,v in [('ascii','Liberation Sans'),('hAnsi','Liberation Sans'),('eastAsia','Droid Sans Fallback'),('cs','Liberation Sans')]:rf.set(qn('w:'+k),v)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+attr),None)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN');rp.append(lang)

fontstyle(doc.styles['Normal'],11)
n=doc.styles['Normal'].paragraph_format;n.space_after=Pt(5);n.line_spacing=1.13;n.widow_control=True
for name,size in [('Title',22),('Subtitle',12),('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
    st=doc.styles[name];fontstyle(st,size,name!='Subtitle');st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(13 if name.startswith('Heading') else 3)
    st.paragraph_format.space_after=Pt(7)
    pp=st.element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
for name,size in [('Caption',9.5),('Header',8.5),('Footer',8.5)]:fontstyle(doc.styles[name],size)
for name in ['Body Small','Quoted Comment','English Draft','Table Body','Table Header','Evidence Note']:
    if name not in doc.styles:doc.styles.add_style(name,1)
    st=doc.styles[name];st.base_style=doc.styles['Normal'];fontstyle(st,10 if name in ['Body Small','Table Body','Table Header','Evidence Note'] else 10.5,name=='Table Header')
    st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.1
doc.styles['Quoted Comment'].paragraph_format.left_indent=Inches(.16)
doc.styles['Quoted Comment'].paragraph_format.right_indent=Inches(.08)
doc.styles['English Draft'].paragraph_format.left_indent=Inches(.16)
doc.styles['Evidence Note'].font.color.rgb=RGBColor.from_string('404040')

# No theme borders or decorative title treatment.
for st in doc.styles:
    if st.type==1:
        pp=st.element.find(qn('w:pPr'))
        if pp is not None:
            for e in list(pp.findall(qn('w:pBdr'))):pp.remove(e)

settings=doc.settings.element
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)

hp=s.header.paragraphs[0];hp.text='PCOS  批注审阅与执行方案';hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('20261030   |   基于Human Reproduction v5   |   ')
for typ,txt in [('begin',None),(None,' PAGE '),('end',None)]:
    r=OxmlElement('w:r')
    if typ:
        v=OxmlElement('w:fldChar');v.set(qn('w:fldCharType'),typ)
    else:
        v=OxmlElement('w:instrText');v.set(qn('xml:space'),'preserve');v.text=txt
    r.append(v);fp._p.append(r)

def p(text='',style=None,lead=None):
    q=doc.add_paragraph(style=style)
    if lead:q.add_run(lead+' ').bold=True
    q.add_run(text.replace('ℛ','花体R'))
    return q

def heading(text,level=1,break_before=False):
    q=doc.add_heading(text,level)
    if break_before:q.paragraph_format.page_break_before=True
    return q

def hyperlink(q,label,url):
    rid=q.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'174A73');rp.append(color)
    under=OxmlElement('w:u');under.set(qn('w:val'),'single');rp.append(under)
    r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);q._p.append(h)

def add_table(headers,rows,widths=None):
    widths=widths or [6.8/len(headers)]*len(headers)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    margins=OxmlElement('w:tblCellMar')
    for key,val in [('top',80),('bottom',80),('left',100),('right',100)]:
        e=OxmlElement('w:'+key);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for ri,vals in enumerate([headers]+rows):
        row=t.rows[0] if ri==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
        for i,(cell,txt,w) in enumerate(zip(row.cells,vals,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            q=cell.paragraphs[0];q.style=doc.styles['Table Header' if ri==0 else 'Table Body'];q.add_run(str(txt))
            q.paragraph_format.space_after=Pt(1);q.paragraph_format.space_before=Pt(1)
            if len(str(txt))<18 and (i==0 or ri==0):q.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EC');cell._tc.get_or_add_tcPr().append(shade)
    q=doc.add_paragraph();q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(0);q.paragraph_format.line_spacing=Pt(2);q.add_run().font.size=Pt(2)
    return t

# Native Office Math building blocks, using real subscripts, fractions and sums.
def mr(txt):
    r=OxmlElement('m:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Latin Modern Math');f.set(qn('w:hAnsi'),'Latin Modern Math');pr.append(f);r.append(pr)
    t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=txt;r.append(t);return r
def parts(x):return [mr(x)] if isinstance(x,str) else (x if isinstance(x,list) else [x])
def wrap(tag,x):
    e=OxmlElement('m:'+tag)
    for z in parts(x):e.append(z)
    return e
def sub(base,idx):
    e=OxmlElement('m:sSub');e.append(wrap('e',base));e.append(wrap('sub',idx));return e
def sup(base,idx):
    e=OxmlElement('m:sSup');e.append(wrap('e',base));e.append(wrap('sup',idx));return e
def frac(a,b):
    e=OxmlElement('m:f');e.append(wrap('num',a));e.append(wrap('den',b));return e
def sm(idx,expr):
    e=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    for tag,val in [('chr','∑'),('limLoc','subSup'),('supHide','1')]:
        n=OxmlElement('m:'+tag);n.set(qn('m:val'),val);pr.append(n)
    e.append(pr);e.append(wrap('sub',idx));e.append(wrap('sup',''));e.append(wrap('e',expr));return e
def bar(x):
    e=OxmlElement('m:bar');pr=OxmlElement('m:barPr');pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');pr.append(pos);e.append(pr);e.append(wrap('e',x));return e
def eq(expr):
    q=doc.add_paragraph();q.paragraph_format.space_before=Pt(5);q.paragraph_format.space_after=Pt(7);q.paragraph_format.line_spacing=1.0
    om=OxmlElement('m:oMathPara');pr=OxmlElement('m:oMathParaPr');jc=OxmlElement('m:jc');jc.set(qn('m:val'),'center');pr.append(jc);om.append(pr)
    math=OxmlElement('m:oMath')
    for z in parts(expr):math.append(z)
    om.append(math);q._p.append(om)

def math_block(name):
    w=lambda:sub('w','a')
    rate=lambda idx:sub('r',idx)
    pop=lambda idx:sub('P',idx)
    share=lambda idx:sub('p',idx)
    if name=='rates':eq([rate('c,a,t,k'),mr(' = '),sup('10','5'),frac(sub('N','c,a,t,k'),pop('c,a,t'))])
    elif name=='shares':eq([share('c,a,t'),mr(' = '),frac(pop('c,a,t'),sm('j∈ℛ',pop('j,a,t'))),mr('     '),w(),mr(' = '),frac(sub('u','a'),sm('b=15–19,…,45–49',sub('u','b')))])
    elif name=='china_ref':
        eq([sub('A','CHN,t,k'),mr(' = '),sm('a',[w(),rate('CHN,a,t,k')])])
        eq([sub('A','R,t,k'),mr(' = '),sm('a',[w(),sm('c∈ℛ',[share('c,a,t'),rate('c,a,t,k')])])])
    elif name=='gap':eq([sub('D','t,k'),mr(' = '),sub('A','CHN,t,k'),mr(' − '),sub('A','R,t,k')])
    elif name=='decomp':eq([sub('ΔD','k;0,T'),mr(' = '),sub('D','T,k'),mr(' − '),sub('D','0,k'),mr(' = '),sub('X','k'),mr(' + '),sub('Q','k'),mr(' + '),sub('S','k')])
    elif name=='x':eq([sub('X','k'),mr(' = '),sm('a',[w(),mr('('),rate('CHN,a,T,k'),mr(' − '),rate('CHN,a,0,k'),mr(')')])])
    elif name=='q':eq([sub('Q','k'),mr(' = −'),sm('a',[w(),sm('c∈ℛ',[bar(share('c,a')),mr('('),rate('c,a,T,k'),mr(' − '),rate('c,a,0,k'),mr(')')])])])
    elif name=='s':eq([sub('S','k'),mr(' = −'),sm('a',[w(),sm('c∈ℛ',[bar(rate('c,a,k')),mr('('),share('c,a,T'),mr(' − '),share('c,a,0'),mr(')')])])])
    elif name=='bars':
        eq([bar(share('c,a')),mr(' = '),frac([share('c,a,0'),mr(' + '),share('c,a,T')],'2'),mr('     '),bar(rate('c,a,k')),mr(' = '),frac([rate('c,a,0,k'),mr(' + '),rate('c,a,T,k')],'2')])
    elif name=='cancellation':
        eq([sub('M','t'),mr(' = |'),sub('D','t,primary'),mr('| + |'),sub('D','t,secondary'),mr('|')])
        eq([sub('K','t'),mr(' = 1 − '),frac([mr('|'),sub('D','t,primary'),mr(' + '),sub('D','t,secondary'),mr('|')],sub('M','t')),mr('     ('),sub('M','t'),mr(' > 0)')])
    elif name=='centred':
        eq([sub('μ','a,k'),mr(' = '),sm('c∈ℛ',[bar(share('c,a')),bar(rate('c,a,k'))])])
        eq([sup(sub('S','c,k'),'∗'),mr(' = −'),sm('a',[w(),mr('('),bar(rate('c,a,k')),mr(' − '),sub('μ','a,k'),mr(')('),share('c,a,T'),mr(' − '),share('c,a,0'),mr(')')])])

def render_section(section,break_before=False):
    heading(section['title'],1,break_before)
    for block in section['blocks']:
        kind=block['type']
        if kind=='p':p(block['text'])
        elif kind=='h2':heading(block['text'],2)
        elif kind=='table':add_table(block['headers'],block['rows'],block.get('widths'))
        elif kind=='equation':math_block(block['name'])

p(DATA['title'],'Title')
p(DATA['subtitle'],'Subtitle')
p(DATA['date'],'Body Small')
p(DATA['execution_status'],'Body Small')
render_section(DATA['front_sections'][0],False)
render_section(DATA['front_sections'][1])

heading('3 全部17条批注的逐条修改方案',1,True)
p('每条按原始ID定位，保留批注原话，并给出现状、具体行动、数据需求、候选英文、验收标准和证据。锚定文字仅为旧稿定位摘录；执行对象为当前最新稿。英文句子是候选文本，涉及已经完成某项分析或已经提供资料的表述，须在对应任务真正验收后使用。')
add_table(['序号与ID','批注主题','v5状态','任务'],[[f"{c['ordinal']:02d}  ID{c['id']}",c['topic'],c['v5_status'],c['task']] for c in DATA['comments']],[1.0,3.2,1.1,1.5])
for c in DATA['comments']:
    heading(f"{c['ordinal']:02d}  ID{c['id']}  {c['topic']}",2)
    meta=p(f"v5状态  {c['v5_status']}    优先级  {c['priority']}    执行任务  {c['task']}",'Body Small')
    meta.paragraph_format.keep_with_next=True
    p(c['original'],'Quoted Comment',lead='原批注')
    anchor=c['anchor'];anchor=anchor[:180]+'…' if len(anchor)>180 else anchor
    p(anchor,'Evidence Note',lead='旧稿定位摘录')
    p(c['diagnosis'],lead='判断与现状')
    ev=next(x for x in CROSS['comments'] if x['comment_id']==c['id']).get('evidence',[])
    if ev:
        x=ev[0];quote=x.get('quote','')
        if quote and '\\' not in quote and '$' not in quote:
            p('“'+quote+'”','Evidence Note',lead='v5实际文本举证')
    for i,action in enumerate(c['actions'],1):p(f'{i}. {action}',lead='具体修改' if i==1 else None)
    p(c['data'],lead='数据与重跑')
    p(c['english_note'],'Evidence Note')
    p(c['english'],'English Draft',lead='英文示例')
    acceptance=p(c['accept'],lead='验收标准');acceptance.paragraph_format.keep_with_next=True
    p(' '.join('['+x+']' for x in c['refs'].split()),'Evidence Note',lead='证据')

for section in DATA['after_sections']:render_section(section,break_before=section['title'].startswith(('8 ','9 ')))
for r in DATA['refs']:
    q=p(r['id']+'  '+r['citation'],'Body Small');q.paragraph_format.keep_with_next=True
    q=p('用于 '+r['use']+'  获取层级 '+r['access']+'。','Body Small')
    q.paragraph_format.keep_with_next=True
    q=p(style='Body Small');hyperlink(q,'原始论文或官方来源',r['url'])
    if r.get('additional_url'):
        q.add_run('   ');hyperlink(q,'期刊正式版PDF',r['additional_url'])

cp=doc.core_properties
cp.title=DATA['title'];cp.subject='17条PCOS论文批注的逐条审阅与可执行修订计划'
cp.author='PCOS研究协作';cp.keywords='PCOS; GBD 2023; Human Reproduction; Codex; 批注'
cp.comments='本文件为审阅与执行方案，不表示新增科研分析已经完成。'
doc.save(OUT)
print(str(OUT))
print('bytes',OUT.stat().st_size,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'references',len(DATA['refs']))
