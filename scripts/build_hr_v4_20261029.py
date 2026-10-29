#!/usr/bin/env python3
"""Dated Human Reproduction revision; reuse audited aggregate estimates, never raw data.

Run with bundled Python (python-docx/Pillow); tectonic and Poppler on PATH.
This script does not submit, send mail, stage Git or synchronize cloud storage.
"""
import argparse, hashlib, json, re, shutil, subprocess, zipfile
from pathlib import Path
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
OLD = ROOT / 'PCOS_Revision_2026-10-29_v3'
NEW = ROOT / 'PCOS_HumanReproduction_v4_2026-10-29'
TITLE = 'Reference-country composition and age-standardized PCOS-attributed infertility contrasts: China and 18 G20 countries, 1990--2023'
RUNNING = 'PCOS infertility and reference composition'
ABSTRACT = [
('STUDY QUESTION', 'How much of the 1990--2023 change in the China-minus-reference age-standardized prevalence difference for total PCOS-attributed infertility is allocated to changing within-age country shares in a contemporaneously population-pooled 18-country reference?'),
('SUMMARY ANSWER', r'Changing reference-country shares contributed +88.65 per 100,000 to the +151.54 change in the modeled difference, equivalent to 58.5\% of the signed net change.'),
('WHAT IS KNOWN ALREADY', 'GBD studies describe PCOS-attributed infertility burden across countries and time. Fixing age-standardization weights does not also fix country population shares within each age group of a multinational reference.'),
('STUDY DESIGN, SIZE, DURATION', 'This exploratory ecological secondary analysis used GBD 2023 estimates for China and 18 other sovereign G20 countries, seven five-year female age groups spanning 15--49 years, and 34 years, 1990--2023. No individual participants were recruited.'),
('PARTICIPANTS/MATERIALS, SETTING, METHODS', 'Matched GBD prevalence counts and female populations yielded age-specific rates per 100,000 females. Fixed age weights were normalized over ages 15--49. The primary estimand was the change in the China-minus-contemporaneous-reference difference for total PCOS-attributed infertility. An established symmetric identity separated Chinese rate, reference-country rate and reference-share components. UN WPP 2024 replaced reference shares only. Subtypes, fixed-1990 shares, equal-country weights and deterministic sensitivity scenarios addressed secondary questions.'),
('MAIN RESULTS AND THE ROLE OF CHANCE', r"China's total rate increased from 333.56 to 605.08 and the reference rate from 606.48 to 726.45 per 100,000. The difference moved from $-272.92$ to $-121.38$: Chinese rates, reference rates and shares contributed $+271.51-208.62+88.65=+151.54$ (rounded). UN shares gave +86.76 for the share component (57.6\% of +150.69). For secondary infertility in 2023, the difference was +2.16 under contemporaneous GBD shares, $-47.99$ under fixed-1990 shares and $-127.89$ under equal-country weights. Jointly aligned GBD posterior draws were unavailable; no derived uncertainty intervals or significance tests are reported."),
('LIMITATIONS, REASONS FOR CAUTION', r'These empirical quantities are conditional on GBD model point estimates and rounded published age weights. The uncertainty of component magnitudes and near-zero contrast signs remains unquantified. Scenario ranges are not uncertainty intervals. The 58.5\% ratio uses a signed net denominator with cancellation, rather than a fraction of women or disease. The G20 reference is economically selected and does not represent a clinical control group or the world.'),
('WIDER IMPLICATIONS OF THE FINDINGS', 'Researchers monitoring national reproductive-health estimates against pooled international populations should state the comparator weights and report absolute rate and composition components. The mathematical persistence of within-age country composition after age standardization is distinct from its estimated magnitude in these data. Benchmark-dependent contrasts do not establish improved care, causal effects or individual infertility risk.'),
('FUNDING', 'XX: verified funding sources and funder roles to be completed by the authors at submission.'),
('DISCLOSURES', 'XX: author-specific disclosures to be completed at submission.'),
('TRIAL REGISTRATION NUMBER', 'Not applicable; this is a secondary analysis of aggregate estimates, not a clinical trial.'),
]

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def write(p, text): p.parent.mkdir(parents=True, exist_ok=True); p.write_text(text)
def replace_once(s, a, b):
    assert s.count(a) == 1, (a[:100], s.count(a))
    return s.replace(a, b, 1)

def build():
    assert not NEW.exists(), 'Keep earlier v4 build; choose a new dated path instead of overwriting.'
    (NEW/'manuscript').mkdir(parents=True)
    for name in ['tables', 'figures']:
        shutil.copytree(OLD/'manuscript'/name, NEW/'manuscript'/name)
    baseline = {str(p.relative_to(OLD)): sha(p) for p in OLD.rglob('*') if p.is_file()}
    write(NEW/'docs/V3_PRESERVATION_BASELINE.json', json.dumps(baseline, indent=2))
    main = (OLD/'manuscript/main.tex').read_text()
    main = main.replace('% EDITORIAL_REVISION_V3_2026_10_29', '% HUMAN_REPRODUCTION_V4_2026_10_29')
    main = main.replace(r'\usepackage[hidelinks]{hyperref}', r'\usepackage[authoryear,round]{natbib}'+'\n'+r'\usepackage[hidelinks]{hyperref}')
    main = main.replace(r'\fancyhead[L]{\small GBD 2023: infertility contrasts}\fancyhead[R]{\small Reproductive Health draft}\fancyfoot[C]{\thepage}', r'\fancyhead[L]{\small '+RUNNING+r'}\fancyhead[R]{\thepage}')
    main = re.sub(r'\\title\{[^\n]*\}', r'\\title{'+TITLE+'}', main, count=1)
    main = main.replace(r'\date{XX Month 2026}', r'\date{29 October 2026}')
    main = main.replace(r'\maketitle\n', r'\maketitle\thispagestyle{fancy}\n') if False else main.replace('\\maketitle\n', '\\maketitle\\thispagestyle{fancy}\n')
    abstract = '\\begin{abstract}\n'+'\n\n'.join(r'\noindent\textbf{'+h+':} '+t for h,t in ABSTRACT)+'\n\\end{abstract}'
    start, end = main.index(r'\begin{abstract}'), main.index(r'\section{Background}')
    main = main[:start]+abstract+'\n\n'+r'\textbf{Keywords:} Polycystic Ovary Syndrome; Infertility, Female; Prevalence; Age Distribution; Population Dynamics.'+'\n\n'+main[end:]
    start, end = main.index(r'\section{Background}'), main.index(r'\section{Methods}')
    intro = r'''\section{Introduction}
Polycystic ovary syndrome (PCOS) is relevant to reproductive health and individualized care \citep{pcos}. Infertility attributed to PCOS in the Global Burden of Disease (GBD) study is a modeled cause-specific impairment outcome, distinct from general PCOS prevalence and infertility risk among couples trying to conceive \citep{gbd}. A practical use of these estimates is to monitor a country's trajectory against a multinational population. Such monitoring requires a definition of whose experience the reference represents at each date.

Liu and colleagues described GBD 2019 PCOS-attributed infertility prevalence, disability and trends at global, regional and national levels \citep{liu_pcos2024}. Liang and colleagues used GBD 2021 and decomposed burden changes into population growth, ageing and epidemiological components \citep{liang_burden2025}. Another analysis examined China and worldwide trends using GBD 2021 \citep{liang_trend2026}. These objectives differ from accounting for a national standardized-rate contrast against an explicitly constructed 18-country pool (Supplementary Table~S9). A count decomposition includes population size and age structure, whereas our contrast decomposition holds the age standard fixed and isolates changing country shares within each age. We do not infer that these earlier studies used our pooling rule or made an erroneous comparison.

A contemporary population-pooled reference is useful when the question concerns the population mix actually represented at each date. Yet the same fixed list of countries can represent a different within-age country mix over time. For example, an increased share for a country with a lower age-specific rate can lower the pooled reference even if each country's rate is unchanged. Fixed age-standardization weights do not remove this second level of composition. A narrowing national gap can therefore coexist with increasing national and reference rates; it cannot by itself establish improved reproductive health.

Our primary objective was to quantify the 1990--2023 change in China's total PCOS-attributed infertility standardized prevalence difference against a contemporaneously population-pooled reference of 18 other sovereign G20 countries, and its symmetric allocation to Chinese rates, reference-country rates and within-age reference-country shares. We apply established Kitagawa-type accounting rather than propose a new method \citep{kitagawa}. Secondary objectives assess demographic-source sensitivity, infertility subtypes and alternative benchmark questions. The fixed G20 country set is interpretable for monitoring but is neither globally representative, a matched clinical control nor a normative treatment target.

'''
    main = main[:start]+intro+main[end:]
    main = main.replace(r'\section{Methods}', r'\section{Materials and methods}')
    main = replace_once(main, 'The statistical analysis was rerun on 29 October 2026; a separate editorial revision on the same date independently audited the source-to-result evidence chain without replacing the archived estimates. ', '')
    main = main.replace('Additional file 1', 'Supplementary data').replace('Additional file 2', 'the reproducibility archive')
    main = main.replace(r'\subsection{Age standardization and the changing reference}', r'''\subsection{Primary estimand and age-standardized reference}
The primary estimand is $\Delta D_{\mathrm{total}}=D_{2023,\mathrm{total}}-D_{1990,\mathrm{total}}$ under contemporaneous GBD reference-country population shares. The population is females aged 15--49 in China and the fixed 18-country reference; the outcome is modeled total PCOS-attributed infertility prevalence, with all females as denominator. The contrast is a difference per 100,000 after restricted-age standardization. Its symmetric components are descriptive allocations of that same change. Subtype contrasts and the UN, fixed-1990 and equal-country specifications are secondary; they do not redefine the primary target.
''')
    main = replace_once(main, 'Thus fixing the country list and the age standard does not fix reference composition. Reference rates were pooled within age before standardization.', 'Thus fixing the country list and the age standard does not fix reference composition. This follows from the pooling equation independently of any GBD estimate; whether the resulting effect is large in these data is an empirical question. Reference rates were pooled within age before standardization.')
    main = replace_once(main, 'The identity is exact before rounding and additive across ages and infertility types.', r'''The identity is exact before rounding and additive across ages and infertility types. Algebraically, $S=0$ if all within-age country shares are unchanged. It is also zero if the mean of each country's two endpoint rates is identical across countries within each age, because share changes sum to zero. These are properties of the identity, not empirical discoveries, statistical null tests or guarantees of a nonzero component in other settings.''')
    main = replace_once(main, 'The choice depends on the intended comparison; none is intrinsically the correct target.', 'The contemporary GBD pool is primary because it matches our monitoring question. The other definitions deliberately change the reference population and therefore the estimand; none is intrinsically the correct target for all questions.')
    # Remove production chronology from scientific methods.
    start = main.index('These exploratory analyses, including')
    end = main.index(r'\subsection{Use of artificial intelligence}', start)
    main = main[:start]+r'''These exploratory analyses were not externally preregistered. Core estimates were independently recomputed in Python 3.12 and base R 4.6 and compared numerically. Code, aggregate results, input specifications and a reproducibility entry point are available in the versioned archive \citep{repository}. Computational checks establish implementation consistency, not independent disease-model validity. The supplement maps applicable GATHER items and identifies incomplete uncertainty reporting \citep{gather}.

\subsection{Ethical approval}
Public aggregate GBD estimates and UN demographic data were analysed without individual participant records. XX: institutional ethics-review applicability or exemption determination to be confirmed by the responsible author. No ethics approval number is asserted.

'''+main[end:]
    # Endometriosis retained entirely in supplement, keeping article focused.
    start = main.index(r'\subsection{Brief contrasting cause: endometriosis}')
    end = main.index(r'\section{Discussion}', start)
    main = main[:start]+main[end:]
    main = re.sub(r' Similarly, the short endometriosis example.*?endpoint\.', '', main)
    main = main.replace("The revised focused analysis omits the earlier draft's ancillary trend-fitting and forecasting modules because they do not resolve the benchmark question. ", '')
    # Stronger comparison and practical reporting implications.
    anchor = r'\section{Discussion}'
    main = replace_once(main, anchor, anchor+r'''
The primary comparison concerns the changing population represented by a pooled benchmark. With fixed age weights, the modeled total China-minus-reference gap moved upward by +151.54 per 100,000; changing reference-country shares contributed +88.65. This empirical magnitude is conditional on GBD point estimates, whereas the persistence of within-age country weighting after age standardization follows directly from the definition. Similar components under UN shares assess demographic-source sensitivity and do not validate the underlying infertility estimates.

Relative to the geographical burden and trend questions studied by Liu and colleagues and Liang and colleagues, our added information is an explicit audit of reference construction \citep{liu_pcos2024,liang_burden2025,liang_trend2026}. We do not reestimate their results or reinterpret country-specific age-standardized rates as pooled comparisons. The methodological contribution is an application of established symmetric accounting to this reproductive-health monitoring question, not a new estimator, a new clinical association or a comprehensive systematic review of prior work.

For practical reporting, investigators should define the population represented by a comparator before examining its direction of change. If monitoring against the contemporary pooled population, report $X$, $Q$ and $S$ in rate units alongside the net difference. If holding the historical country mix is the intended question, specify that fixed-share target directly. An equal-country target gives a small and a large country equal influence within age. Showing these targets side by side reveals how a near-zero contrast depends on the question, rather than selecting whichever benchmark yields a preferred sign.
''')
    start=main.index('The main finding is a measurement issue:')
    end=main.index('The absolute identity is more informative',start)
    main=main[:start]+main[end:]
    # Replace endmatter with journal order. Preserve intentional submission XX fields.
    start = main.index(r'\section{Conclusions}')
    end = main.index(r'{\small\setstretch{1.2}', start)
    main = main[:start]+r'''\section*{Authors' roles}
XX: verified author-by-author contributions to design, execution, analysis, drafting and critical discussion, final approval and accountability to be supplied at submission.

\section*{Acknowledgements}
We acknowledge the GBD data providers and collaborators and the United Nations Population Division. XX: additional acknowledgements, if applicable.

\section*{Funding}
XX: verified sponsors, grants and funder roles, or an author-approved no-specific-funding statement, to be supplied at submission.

\section*{Disclosures}
XX: author-specific financial, personal and professional disclosures to be supplied at submission.

\section*{Data availability}
Code, aggregate derived results, input specifications and checksums are available in the public repository \citep{repository}. The accompanying reproducibility archive provides the executable Python and R workflow. Original licensed GBD CSVs must be obtained through the GBD Results Tool \citep{results} using the stated dimensions; they are retained in the private research archive and excluded from the public package, together with country--age source-cell extracts. The UN subset is available with source attribution \citep{wpp2024}. No hospital data were used. The immutable publication revision and byte-verification receipt accompany this manuscript.

\section*{Supplementary data}
Supplementary methods, eight numerical tables, seven figures, a literature/estimand comparison table (S9), and a GATHER reporting map accompany the article. The reproducibility archive contains code, aggregate results, source specifications and a README; licensed GBD originals and restricted source cells are excluded.

'''+main[end:]
    main = main.replace(r'\cite{', r'\citep{')
    # Author-year labels and alphabetical reference order required by HR.
    labels = {'chen':'Chen et al.(2026)', 'gather':'Stevens et al.(2016)', 'gbd':'GBD 2023 Collaborators(2025)', 'liu_pcos2024':'Liu et al.(2024)', 'liang_burden2025':'Liang et al.(2025)', 'liang_trend2026':'Liang(2026)', 'pcos':'Teede et al.(2023)', 'repository':'PCOS research workspace(2026)', 'results':'Institute for Health Metrics and Evaluation(2026)', 'wpp2024':'United Nations(2024)', 'xu':'Xu et al.(2025)', 'kitagawa':'Kitagawa(1955)'}
    # Xu no longer directly used: retain context citation in sensitivity methods.
    main = main.replace('Detailed endometriosis terminal-year checks are confined to Supplementary data.', r'Detailed endometriosis terminal-year checks, a secondary cause contrast related to broader attribution studies \citep{xu}, are confined to Supplementary data.')
    bs, be = main.index(r'\begin{thebibliography}{99}'), main.index(r'\end{thebibliography}')
    items = re.findall(r'\\bibitem\{([^}]+)\}(.*?)(?=\\bibitem|$)', main[bs+len(r'\begin{thebibliography}{99}'):be], re.S)
    entries = {key:r'\bibitem['+labels[key]+']{'+key+'}'+body.strip()+'\n' for key,body in items}
    order = ['chen','gbd','results','kitagawa','liang_trend2026','liang_burden2025','liu_pcos2024','repository','gather','pcos','wpp2024','xu']
    main = main[:bs]+r'\begin{thebibliography}{99}'+'\n'+''.join(entries[k] for k in order)+main[be:]
    # Alt text directly below each legend, no invented error bars.
    alts = [
      'Two-panel line plot. China and the population-pooled reference both increase from 1990 to 2023; the China-minus-reference difference remains negative and narrows from minus 272.92 to minus 121.38 per 100,000.',
      'Primary and secondary PCOS-attributed infertility trajectories. The primary contrast remains negative; the secondary contrast approaches zero and is plus 2.16 per 100,000 in 2023 under contemporary GBD reference shares.',
      'Bars separate Chinese rate, reference rate and reference share components by infertility type, with diamonds for net changes. For total infertility the components are plus 271.51, minus 208.62 and plus 88.65 per 100,000.',
      'Annual differences and endpoint components under GBD and UN shares are close. For total infertility, the share component is plus 88.65 with GBD and plus 86.76 with UN; disease rates are held fixed.',
      'Benchmark bars for secondary infertility in 2023. The GBD difference is plus 2.16 with current shares, minus 47.99 with 1990 shares and minus 127.89 with equal-country weights; UN scenarios show the same sign pattern.'
    ]
    for i, alt in enumerate(alts):
        label = ['trends','subtypes','decomp','source','bench'][i]
        main = main.replace(r'\label{fig:'+label+r'}\end{figure}', r'\label{fig:'+label+'}\n'+r'\caption*{\textbf{Alt text:} '+alt+r'}\end{figure}')
    main=main.replace(r'\begin{abstract}', '\\begingroup\\sloppy\n'+r'\begin{abstract}').replace(r'\end{abstract}', r'\end{abstract}'+'\n'+r'\endgroup')
    write(NEW/'manuscript/main.tex', main)
    supp = (OLD/'manuscript/supplement.tex').read_text()
    supp = supp.replace('% INTEGRATED_REVISION_V2','% HUMAN_REPRODUCTION_V4_SUPPLEMENT')
    supp = supp.replace(r'\fancyhead[R]{\small Additional file 1}\fancyfoot[C]{\thepage}',r'\fancyhead[R]{\thepage}')
    supp = supp.replace('Additional file 1:', 'Supplementary data:').replace('Companion to the revised comparator manuscript', 'Companion to Human Reproduction v4')
    a, b = supp.index('This supplement accompanies'), supp.index('The GBD inputs comprise')
    supp = supp[:a]+'''This supplement documents the secondary computations accompanying the main article. The primary estimand is the 1990--2023 change in the total PCOS-attributed infertility difference under contemporary GBD reference shares. Numerical assets are reused from the independently verified analysis. This revision changes scientific framing and reporting, without introducing new observations or reestimating the upstream GBD model.\n\n'''+supp[b:]
    lit = r'''
\clearpage
\setcounter{table}{8}
\section*{Prior questions and the present estimand}
This targeted comparison supports scientific positioning; it is not an exhaustive systematic review or proof that no prior analysis has used similar accounting. Bibliographic details and source URLs are in the companion evidence audit. No pooled-comparator error is attributed to a paper simply because it reports national age-standardized rates.
\small
\begin{longtable}{p{.17\textwidth}p{.34\textwidth}p{.39\textwidth}}
\caption{Questions addressed by selected prior PCOS-attributed infertility analyses and the present study.}\\
\toprule Study & Verified question or analysis & Relation to the present primary estimand\\\midrule\endhead
Liu et al., 2024 & GBD 2019 global, regional and national prevalence, YLD and trends, 1990--2019. Published Methods describe direct age standardization. & Provides disease-burden context. It does not establish that a China-versus-18-country pool with our within-age weights was used. We do not reinterpret its individual national comparisons.\\
Liang et al., 2025 & GBD 2021 burden and trends; the Methods describe a Das Gupta allocation to population growth, ageing and epidemiological factors. & A burden/count change and a national restricted-age standardized-rate difference against a changing pool are different targets. Our share term concerns countries within each age, not changing age structure.\\
Liang, 2026 & Indexed abstract: GBD 2021 China and worldwide PCOS/related-infertility trends, including Joinpoint and ARIMA. Full comparator construction was not independently verified. & Supports a China/worldwide monitoring use case. No claim about an undocumented pooling formula is made; conclusions here do not test that paper's estimates.\\
Present study & GBD 2023 China-minus-contemporaneous-18-country pooled total-infertility difference, 1990--2023; fixed age weights. & Primary change and established symmetric allocation to national rates, reference rates and within-age country shares. UN shares test demographic-source sensitivity; fixed-1990/equal-country weights change the target.\\\bottomrule
\end{longtable}\normalsize
'''
    supp = supp.replace(r'\end{document}', lit+'\n'+r'\end{document}')
    # Restore references to public archive, not missing v4 code/data folders.
    supp = supp.replace(r'\texttt{code/revision\_analysis.py}', r'\texttt{PCOS\_Revision\_2026-10-29/code/revision\_analysis.py}')
    supp=supp.replace(r'\begin{document}',r'\hypersetup{hypertexnames=false}'+'\n'+r'\begin{document}')
    supp=supp.replace(r'\maketitle',r'\maketitle\thispagestyle{fancy}').replace('Additional file 2','Reproducibility archive').replace('Main Declarations:','Main Funding section:')
    write(NEW/'manuscript/supplement.tex', supp)
    cover = r'''\documentclass[11pt,a4paper]{article}
\usepackage[margin=27mm]{geometry}
\usepackage[T1]{fontenc}
\usepackage{mathpazo,setspace}
\usepackage[hidelinks]{hyperref}
\setstretch{1.08}\setlength{\parindent}{0pt}\setlength{\parskip}{9pt}
\pagestyle{empty}
\begin{document}
29 October 2026\\
Editor-in-Chief\\
\emph{Human Reproduction}

Dear Editor-in-Chief,

Please consider our manuscript, ``'''+TITLE+r'''", as an Original Article.

The scientific question is how much of the changing China-minus-international-reference difference in modeled PCOS-attributed infertility is allocated to changing country population shares within age, even when age-standardization weights remain fixed. This is relevant to readers using international burden estimates to monitor reproductive health: the reference population represented at each date is part of the question being answered.

In GBD 2023 estimates for females aged 15--49, the total China-minus-18-country difference changed by +151.54 per 100,000 between 1990 and 2023. Established symmetric accounting allocated +88.65 to reference-country composition, alongside +271.51 from Chinese rates and $-208.62$ from reference-country rates. Replacing reference shares with UN WPP 2024 gave a similar +86.76 composition component. Alternative benchmark definitions changed the sign of a near-zero secondary-infertility contrast in 2023. We report absolute components before ratios and distinguish demographic sensitivity from statistical uncertainty.

Liu and colleagues' 2024 \emph{Human Reproduction} paper (doi:10.1093/humrep/dead241) described PCOS-attributed infertility burden and geographical trends. Our manuscript addresses the different question of how an explicitly pooled reference is constructed. Its contribution is a focused reproductive-epidemiology application and reporting implication, not a new decomposition algorithm or a claim that earlier national rate comparisons were erroneous. It does not estimate individual infertility risk, treatment benefit or causal effects.

The work is an exploratory secondary analysis of aggregate estimates. Confirmed jointly aligned GBD posterior draws were unavailable; uncertainty intervals for our custom contrasts and components are therefore not supplied. This limitation is explicit in the abstract and article. Code, aggregate results and a reproducibility archive are provided; licensed GBD originals and restricted source-cell extracts are excluded from the public repository. OpenAI-assisted coding and drafting are disclosed in the manuscript.

XX: Corresponding author to confirm final author approval, authorship order, exclusive consideration, related manuscripts/preprints, funding and disclosures, and institutional ethics-review applicability before submission. These declarations have not been assumed on behalf of the authors.

Yours sincerely,\\
XX Corresponding Author\\
XX Institution; XX email; ORCID: XX
\end{document}
'''
    cover=cover.replace(r'\begin{document}',r'\begin{document}\sloppy').replace(r'\"',"''")
    write(NEW/'manuscript/cover_letter.tex', cover)
    evidence = {
      'checked_on':'2026-10-29','abstract_template_version':'20/05/2026',
      'official_sources':[
        {'url':'https://academic.oup.com/humrep/pages/Original_Article','access':'browser read','scope':'article type routes and ORCID'},
        {'url':'https://academic.oup.com/humrep/pages/Clinical_Studies','access':'browser read','scope':'title<=25 words, running title<=50 chars, <=5 keywords, extended abstract, journal order, author-date references, separate editable Word tables, separate jpg/tif figures'},
        {'url':'https://academic.oup.com/humrep/pages/Manuscript_Preparation','access':'browser read','scope':'double spacing, continuous line numbers, page number top right, cover letter, authorship signatures, disclosures, alt text, data availability'},
        {'url':'https://static.primary.prod.gcms.the-infra.com/static/site/humrep/document/clinical-study-extended-abstract.doc?node=7434962e6537ca0ef7c5','access':'download and textutil read','scope':'11 exact headings; FUNDING and DISCLOSURES separate; no numerical abstract cap found in template'},
        {'url':'https://academic.oup.com/humrep/pages/Submission_Online','access':'read during preceding assessment','scope':'Word-oriented upload guidance; LaTeX conversion not automatic; line art 1200 dpi and colour 300 dpi guidance'},
        {'url':'https://academic.oup.com/humrep/pages/Policies','access':'read during preceding assessment','scope':'ethics/AI/editorial policies; author verification still needed'},
      ],
      'literature':[
       {'doi':'10.1093/humrep/dead241','url':'https://academic.oup.com/humrep/article/39/1/108/7452815','access':'publisher full text via browser','verified':'GBD2019 prevalence/YLD/trends; direct age standardization; no accusation of exact comparator use'},
       {'doi':'10.3389/fpubh.2025.1693486','url':'https://www.frontiersin.org/journals/public-health/articles/10.3389/fpubh.2025.1693486/full','access':'publisher full text','verified':'GBD2021 burden; Methods: Das Gupta population growth, ageing, epidemiological factors'},
       {'doi':'10.1016/j.rbmo.2025.105239','url':'https://pubmed.ncbi.nlm.nih.gov/41456540/','access':'indexed abstract verified in prior audit','verified':'China/worldwide trends and Joinpoint/ARIMA; full comparator definition unverified; do not infer formula'}],
      'limits':'No current numerical main-text word or combined figure/table cap found on the accessible official general/type pages. No invented 3500-word, 600-word or 6-object rule. Recheck live portal at eventual submission; no submission performed.'}
    write(NEW/'docs/HR_REQUIREMENTS_AND_EVIDENCE.json',json.dumps(evidence,ensure_ascii=False,indent=2))
    changes = '''# Human Reproduction v4: scientific positioning and preparation

Prepared 29 October 2026. v3 is preserved byte-for-byte. No submission or email sent.

1. Actual use case and prior evidence: the Introduction defines longitudinal monitoring against a contemporary population-pooled fixed country set. Supplementary Table S9 compares the verified aims of Liu 2024, Liang 2025 and Liang 2026, with full-text/abstract access limits. It does not allege earlier national ASPR comparisons used our pool or were erroneous. No exhaustive priority or new-method claim is made.
2. Primary estimand: change from 1990 to 2023 in total PCOS-attributed infertility China-minus-18-country age-standardized difference under contemporary GBD shares. Subtypes, UN shares and other benchmarks are secondary and explicitly answer other questions. Target population, outcome, denominator, time, comparison and units are stated.
3. Algebra versus evidence: the standardization equation and zero-component conditions are mathematical properties. +88.65 and 58.5% are empirical, conditional GBD point-estimate quantities. The lack of aligned posterior draws, signed-denominator cancellation and limits on near-zero signs are explicit. No intervals were fabricated and no new disease observations were added.

Journal preparation: dated May 2026 extended abstract with separate FUNDING and DISCLOSURES; five keywords; author-date alphabetized references; authors' roles directly after Discussion; separate ethics subsection; Data availability; right-top page numbers; continuous line numbers; main figure Alt text. Existing five numerical figures and three tables retained. Endometriosis is confined to the supplement. Production-history narration removed from main scientific methods.

The article PDF is a review copy with figures/tables embedded for reading. Editable Tables.docx and high-resolution separate figure TIFFs are supplied. TeX is the authoritative editable article; a full Word article has not been produced. The journal's current upload guidance discourages LaTeX because automatic PDF conversion is unavailable. Before eventual submission, the corresponding author should use the portal's supported route or convert and proofread the complete article in Word. This is a remaining packaging step, not a scientific result.

Remaining author tasks (intentionally not blocking this revision): fill XX identity/affiliation/ORCID, author-specific contributions, funding/disclosures; establish institutional ethics-review applicability; obtain all author signatures and final approval; verify exclusivity/related preprints. No approval, absence of conflicts or exclusivity is asserted. Joint posterior uncertainty remains scientifically unavailable; these revisions improve positioning without converting the study into a clinical validation or a guaranteed Q1 acceptance.

No new raw/processed research data: DATA_CATALOG unchanged. All numerical table fragments and figure PDFs are identical to v3. Reproducibility code/results remain in PCOS_Revision_2026-10-29, referenced by the archive. Private licensed raw files remain in the existing archive.
'''
    write(NEW/'docs/CHANGES_AND_PRE_SUBMISSION_ZH.md',changes)
    write(NEW/'README.md','''# Human Reproduction v4 — 29 October 2026

Independent dated manuscript and cover letter. No submission performed. Old v3 retained.

- `manuscript/main.tex`, `supplement.tex`, `cover_letter.tex`: authoritative sources.
- `output/PCOS_HumanReproduction_v4.pdf`, `PCOS_HumanReproduction_v4_Supplement.pdf`, `HumanReproduction_Cover_Letter_v4.pdf`: reading copies.
- `submission/Tables.docx`: separate editable tables; `Cover_Letter.docx`: editable draft; `submission/figures/`: separate 1200-dpi TIFF line-art exports from the verified vector figures.
- `docs/HR_REQUIREMENTS_AND_EVIDENCE.json`, `CHANGES_AND_PRE_SUBMISSION_ZH.md`: source audit and remaining author/portal checks.
- Existing full analysis code/results/input hashes: `../PCOS_Revision_2026-10-29/`; independently audited v3 computations: `../PCOS_Revision_2026-10-29_v3/docs/`.

Build from repository root with bundled Python: `python scripts/build_hr_v4_20261029.py build` (new folder only). Compile existing sources with `... compile`; verify with `... verify`. Tectonic and Poppler are required. Licensed original data are neither copied into this revision nor included in the public bundle. The complete archive includes the reviewed public reproduction paths from the v3 manifest.

The primary analysis remains exploratory and conditional on GBD model point estimates. No joint uncertainty intervals, clinical treatment claims or ethics approval are invented. Intentional XX author declarations must be completed at submission. See the preparation note regarding journal Word/upload compatibility.
''')
    compile_all()
    make_docs(main, cover)

def compile_all():
    out = NEW/'output'; out.mkdir(exist_ok=True)
    for src,dest in [('main','PCOS_HumanReproduction_v4'),('supplement','PCOS_HumanReproduction_v4_Supplement'),('cover_letter','HumanReproduction_Cover_Letter_v4')]:
        subprocess.run(['tectonic','--keep-logs','--outdir',str(out),str(NEW/'manuscript'/f'{src}.tex')],cwd=NEW/'manuscript',check=True)
        shutil.move(out/f'{src}.pdf',out/f'{dest}.pdf')

def tex_plain(s):
    s=s.replace('---','—').replace('--','–').replace(r'\%', '%').replace(r'\_', '_').replace(r'\&','&').replace('~',' ').replace('``','“').replace("''",'”')
    s=re.sub(r'\\(?:emph|textbf|texttt|mathrm)\{([^{}]*)\}',r'\1',s)
    s=re.sub(r'\\href\{[^}]+\}\{([^}]+)\}',r'\1',s)
    s=re.sub(r'\\url\{([^}]+)\}',r'\1',s)
    s=s.replace('$','').replace(r'\Delta','Δ').replace(r'\times','×')
    return re.sub(r'\\[a-zA-Z]+','',s).replace('{','').replace('}','').strip()

def make_docs(main, cover, export_figures=True):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    def newdoc():
        d=Document(); sec=d.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.8)
        sec.left_margin=sec.right_margin=Inches(.8)
        st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(11); st.paragraph_format.space_after=Pt(6)
        for style in ['Title','Heading 1','Heading 2']:
            d.styles[style].font.name='Times New Roman'; d.styles[style].font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        p=sec.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE');p._p.append(f)
        return d
    submission=NEW/'submission'; submission.mkdir(exist_ok=True)
    d=newdoc();d.add_heading('Tables',0)
    for i,match in enumerate(re.finditer(r'\\begin\{table\}.*?\\end\{table\}', main,re.S),1):
        if i>1:d.add_page_break()
        block=match.group(); cap=re.search(r'\\caption\{(.*?)\}\\label',block,re.S).group(1)
        d.add_heading('Table '+str(i)+'. '+tex_plain(cap),1)
        name=re.search(r'\\input\{tables/([^}]+)\}',block).group(1)
        headers={'prevalence.tex':['Type','Year','China','Reference','Difference','Ratio'], 'decomposition.tex':['Type','Shares','China rates X','Ref. rates Q','Ref. shares S','Net ΔD'], 'benchmark.tex':['Population source','Benchmark','Reference rate','China minus reference']}[name]
        lines=(NEW/'manuscript/tables'/name).read_text().strip().splitlines()
        rows=[[tex_plain(c) for c in ln.replace(r'\\','').split('&')] for ln in lines if '&' in ln]
        table=d.add_table(rows=1, cols=len(headers));table.style='Table Grid'
        for c,h in zip(table.rows[0].cells,headers):c.text=h
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            assert len(row)==len(headers),row
            for c,t in zip(table.add_row().cells,row):c.text=t
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:r.font.size=Pt(9)
        note=re.search(r'\\caption\*\{(.*?)\}\s*\\end\{table\}',block,re.S).group(1)
        d.add_paragraph(tex_plain(note))
    d.save(submission/'Tables.docx')
    d=newdoc()
    body=cover.split(r'\begin{document}',1)[1].split(r'\end{document}')[0]
    for para in body.strip().split('\n\n'):
        d.add_paragraph(tex_plain(para.replace(r'\\','\n')))
    d.save(submission/'Cover_Letter.docx')
    # Vector PDF -> true 1200dpi exports; no raster upsampling of existing PNGs.
    figs=submission/'figures';figs.mkdir(exist_ok=True)
    for p in (sorted((NEW/'manuscript/figures').glob('Figure*.pdf')) if export_figures else []):
        subprocess.run(['pdftoppm','-r','1200','-singlefile','-tiff','-tiffcompression','lzw',str(p),str(figs/p.stem)],check=True)

def verify():
    main=(NEW/'manuscript/main.tex').read_text(); supp=(NEW/'manuscript/supplement.tex').read_text()
    base=json.loads((NEW/'docs/V3_PRESERVATION_BASELINE.json').read_text())
    assert all(sha(OLD/p)==h for p,h in base.items()), 'v3 changed'
    assert all(sha(p)==sha(OLD/'manuscript'/p.relative_to(NEW/'manuscript')) for folder in ['tables','figures'] for p in (NEW/'manuscript'/folder).iterdir() if p.is_file())
    assert len(TITLE.replace('--','–').split())<=25 and len(RUNNING)<=50
    for h,t in ABSTRACT:assert r'\textbf{'+h+':}' in main
    assert len(ABSTRACT)==11
    assert main.count('Alt text:')==5
    assert main.count(r'\begin{table}')==3 and main.count(r'\begin{figure}')==5
    assert 'primary estimand' in main.lower() and 'Algebraically' in main
    assert r'\section*{Authors' in main and main.index(r'\section*{Authors')>main.index(r'\section{Discussion}')
    assert 'Reproductive Health draft' not in main and 'editorial revision' not in main.lower()
    assert 'Supplementary Table~S9' in main and r'\setcounter{table}{8}' in supp
    for p in (NEW/'output').glob('*.pdf'):
        # Bundled Poppler exposes pdftoppm/pdfinfo wrappers, but pdftotext may
        # reside only in the native distribution. pypdf avoids a PATH assumption.
        from pypdf import PdfReader
        result='\n'.join(page.extract_text() or '' for page in PdfReader(p).pages)
        assert len(result)>1000 and 'XX' in result
        assert '??' not in result, p
    from docx import Document
    d=Document(NEW/'submission/Tables.docx');assert len(d.tables)==3
    editable_rows_checked=0
    for table,name in zip(d.tables,['prevalence.tex','decomposition.tex','benchmark.tex']):
        source_rows=[[tex_plain(c) for c in ln.replace(r'\\','').split('&')] for ln in (NEW/'manuscript/tables'/name).read_text().strip().splitlines() if '&' in ln]
        assert [[c.text for c in row.cells] for row in table.rows[1:]]==source_rows
        editable_rows_checked+=len(source_rows)
    from PIL import Image
    imeta=[]
    for p in sorted((NEW/'submission/figures').glob('*.tif')):
        with Image.open(p) as im:imeta.append({'name':p.name,'pixels':im.size,'dpi':[float(x) for x in im.info.get('dpi')]});assert im.info.get('dpi')[0]==1200
    checks={'verified_at':datetime.now().astimezone().isoformat(),'title_words':len(TITLE.split()),'running_title_characters':len(RUNNING),'abstract_approx_words':len(tex_plain(' '.join(t for _,t in ABSTRACT)).split()),'keywords':5,'extended_abstract_headings':11,'v3_files_unchanged':len(base),'numerical_assets_unchanged':True,'main_tables':3,'main_figures':5,'main_alt_texts':5,'separate_editable_tables':3,'editable_table_rows_verified':editable_rows_checked,'tiff_exports':imeta,'statistical_analysis_rerun':False,'new_research_data':False,'posterior_intervals_available':False,'submitted':False}
    write(NEW/'docs/DOCUMENT_AND_SCIENCE_QC.json',json.dumps(checks,indent=2))
    print(json.dumps(checks,indent=2))

def package():
    old_manifest=json.loads((OLD/'docs/PUBLIC_ARTIFACT_MANIFEST.json').read_text())['files']
    # Reuse reviewed original public reproduction paths, but not obsolete v3 manuscript PDFs.
    paths=[ROOT/f['path'] for f in old_manifest if f['path'].startswith('PCOS_Revision_2026-10-29/')]
    paths += [p for p in NEW.rglob('*') if p.is_file() and p.suffix not in ['.zip','.log'] and '/rendered/' not in str(p)]
    paths += [Path(__file__).resolve()]
    paths=sorted(set(paths))
    entries=[{'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':sha(p)} for p in paths]
    write(NEW/'docs/PUBLIC_ARTIFACT_MANIFEST.json',json.dumps({'licensed_raw_GBD_included':False,'country_age_source_cells_included':False,'patient_data_included':False,'files':entries},indent=2))
    paths.append(NEW/'docs/PUBLIC_ARTIFACT_MANIFEST.json')
    with zipfile.ZipFile(NEW/'output/PCOS_HumanReproduction_v4_Complete.zip','w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in paths:z.write(p,p.relative_to(ROOT))
    print('Public archive files:',len(paths))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('action',choices=['build','compile','docs','verify','package']);a=parser.parse_args()
    {'build':build,'compile':compile_all,'docs':lambda:make_docs((NEW/'manuscript/main.tex').read_text(),(NEW/'manuscript/cover_letter.tex').read_text(),False),'verify':verify,'package':package}[a.action]()
