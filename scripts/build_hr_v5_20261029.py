#!/usr/bin/env python3
"""Independent HR v5 technical revision. Reuses audited estimates; never edits v4."""
import importlib.util
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
OLD = ROOT/'PCOS_HumanReproduction_v4_2026-10-29'
NEW = ROOT/'PCOS_HumanReproduction_v5_2026-10-29'
PRIVATE = ROOT/'.shared_context_runtime/v5'
SNAPSHOT = '1abf337d9f23cc69292e03cd581b6a1557e6effe'
REPO = 'https://github.com/jiabao1314/PCOS'
STABLE = REPO+'/tree/'+SNAPSHOT+'/PCOS_Revision_2026-10-29'
UN_URL = 'https://population.un.org/wpp/assets/Excel%20Files/1_Indicator%20(Standard)/CSV_FILES/WPP2024_PopulationByAge5GroupSex_Medium.csv.gz'
RUNNING = 'PCOS infertility and reference composition'
AI = ("OpenAI ChatGPT and Codex supported code development, analysis checks, drafting and technical editing under researcher-specified questions and revision instructions. "
      "The recorded model identifier for the documented Codex session on 29 October 2026 was \\texttt{gpt-6.1-sol}; exact model versions for earlier separate ChatGPT sessions could not be independently verified and are not specified. "
      "No artificial observations or posterior draws replaced study data. The authors' final scientific review, verification of AI-supported content and approval remain to be confirmed before submission; human authors retain responsibility for the work.")

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def write(p,s): p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s)
def save(p,d): write(p,json.dumps(d,ensure_ascii=False,indent=2)+'\n')
def module(name,path):
    spec=importlib.util.spec_from_file_location(name,path);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m
def once(s,a,b):
    assert s.count(a)==1,(a[:90],s.count(a));return s.replace(a,b,1)
def marker(format,count):
    flag=PRIVATE/('marker_'+format)
    if flag.exists():return
    base=Path('/Users/wangjiajia/.codex/plugins/cache/openai-primary-runtime')
    skill='documents' if format=='docx' else 'pdf'
    subprocess.run(['/Users/wangjiajia/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node',str(base/skill/'26.1026.10807/skills'/skill/'container_tools/mark_artifact_operation_started.mjs'),'--operation-kind','create','--expected-output-count',str(count),'--output-format',format],check=True)
    flag.write_text('successful marker invocation\n')

def source():
    assert not NEW.exists(),'Preserve existing v5; do not overwrite an earlier revision'
    PRIVATE.mkdir(exist_ok=True)
    baseline={str(p.relative_to(ROOT)):sha(p) for p in OLD.rglob('*') if p.is_file()}
    (NEW/'manuscript').mkdir(parents=True)
    for name in ['tables','figures']:shutil.copytree(OLD/'manuscript'/name,NEW/'manuscript'/name)
    shutil.copytree(OLD/'submission/figures',NEW/'submission/figures')
    save(NEW/'docs/V4_PRESERVATION_BASELINE.json',baseline)
    main=(OLD/'manuscript/main.tex').read_text().replace('HUMAN_REPRODUCTION_V4','HUMAN_REPRODUCTION_V5')
    main=once(main,r'\maketitle\thispagestyle{fancy}',r'\maketitle\thispagestyle{fancy}'+'\n'+r'\noindent\textbf{Running title:} '+RUNNING+'\n')
    start=main.index(r'\subsection{Use of artificial intelligence}')
    end=main.index(r'\section{Results}',start)
    main=main[:start]+r'''\subsection{Use of artificial intelligence}
OpenAI ChatGPT and Codex supported code development, analysis checks, drafting and technical editing. The documented Codex session used the recorded model identifier \texttt{gpt-6.1-sol} on 29 October 2026. Further disclosure and the pending author-review status are given in Acknowledgements.

'''+main[end:]
    main=once(main,'XX: additional acknowledgements, if applicable.','XX: additional acknowledgements, if applicable.\n\n'+AI)
    start=main.index(r'\section*{Data availability}');end=main.index(r'\section*{Supplementary data}',start)
    main=main[:start]+r'''\section*{Data availability}
The executable Python and R workflow, aggregate derived results, query specifications and checksums are available in the immutable public research snapshot \citep{repository}. The snapshot identifies the analysis underlying the present article and is distinct from the manuscript version. Licensed original GBD 2023 extracts were accessed through the GBD Results Tool on 25 October 2026 \citep{results}; the query dimensions, export identifiers and SHA-256 hashes needed to identify them are recorded in the snapshot's input manifest. These originals and country--age source-cell extracts are retained in the private research archive and are excluded from the public repository. UN WPP 2024 population data are identified by the exact source product and file in the reference list \citep{wpp2024}; the attributed 4,522-row subset and its source manifest are included in the immutable snapshot. No hospital data were used.

'''+main[end:]
    substitutions={
        'results':r'[dataset] Institute for Health Metrics and Evaluation. GBD 2023 prevalence and matched female-population extracts for China and 18 other G20 countries, 1990--2023. GBD Results Tool; extracts accessed 25 October 2026. \url{https://vizhub.healthdata.org/gbd-results/}. Release: GBD 2023; export identifiers 1b13e337, 916826fb, 416daf73 and 76c5523c. Query specifications and input hashes are provided in the versioned research snapshot \citep{repository}. Access is subject to IHME terms.',
        'repository':r'[dataset] PCOS research workspace. Analysis code, query manifests and aggregate derived results for GBD 2023 infertility comparisons. GitHub; 2026. Immutable revision '+SNAPSHOT+r'. \url{'+STABLE+r'}. Includes the Python/R workflow, input manifest, aggregate results and attributed UN population subset; restricted GBD originals are excluded.',
        'wpp2024':r'[dataset] United Nations, Department of Economic and Social Affairs, Population Division. World Population Prospects 2024, Online Edition: population by five-year age group and sex, medium variant, midyear. United Nations; 2024. Source file: \texttt{WPP2024\_PopulationByAge5GroupSex\_Medium.csv.gz}. \url{'+UN_URL+r'}. Historical estimates through 2023 only were selected. Creative Commons Attribution 3.0 IGO.'}
    for key,body in substitutions.items():
        pattern=r'(\\bibitem\[[^\]]+\]\{'+key+r'\})(.*?)(?=\n\\bibitem|\n\\end\{thebibliography\})'
        main,n=re.subn(pattern,lambda m:m[1]+body,main,flags=re.S);assert n==1,key
    # A bibliography item must not itself create an additional author-date citation.
    main=main.replace(r'Query specifications and input hashes are provided in the versioned research snapshot \citep{repository}. Access is subject to IHME terms.',r'Query specifications and input hashes are provided in the versioned research snapshot listed below. Access is subject to IHME terms.')
    write(NEW/'manuscript/main.tex',main)
    supp=(OLD/'manuscript/supplement.tex').read_text().replace('HUMAN_REPRODUCTION_V4','HUMAN_REPRODUCTION_V5')
    fulltitle=re.search(r'\\title\{([^\n]+)\}',main)[1]
    supp=once(supp,r'\title{Supplementary data: Reference-country composition and PCOS-attributed infertility}',r'\title{Supplementary data: '+fulltitle+'}')
    supp=once(supp,r'\author{Companion to Human Reproduction v4}',r'\author{Companion supplementary methods and results}')
    supp=once(supp,'Numerical assets are reused from the independently verified analysis. This revision changes scientific framing and reporting, without introducing new observations or reestimating the upstream GBD model.','The analyses use aggregate GBD estimates and matched demographic data; the upstream GBD model was not reestimated.')
    supp=once(supp,'The prior published aggregate tables are used only for post-computation regression checks, not to generate results.','Archived aggregate tables are used only for post-computation regression checks, not to generate results.')
    start=supp.index('The accessible GBD input files provide point estimates and marginal lower/upper limits, without draw columns.')
    end=supp.index('Until official joint alignment is confirmed',start)
    supp=supp[:start]+'''The GBD input files contain point estimates and marginal lower/upper limits, without posterior draw columns. Availability of a provider-confirmed joint posterior array aligned across all dimensions required by this analysis was not established from the study inputs. This does not imply that suitable draws do not exist or cannot be supplied. The present analyses therefore do not quantify joint posterior uncertainty.

'''+supp[end:]
    supp=once(supp,"Author-provided declarations remain XX at the user's request for completion at submission.",'Author-provided declarations remain XX pending author confirmation before submission.')
    supp=once(supp,'Bibliographic details and source URLs are in the companion evidence audit.','Bibliographic details and source URLs are provided in the main article reference list.')
    write(NEW/'manuscript/supplement.tex',supp)
    oldcover=(OLD/'manuscript/cover_letter.tex').read_text()
    head=oldcover.split(r'\begin{document}')[0]
    declaration=re.search(r'XX: Corresponding author to confirm.*?authors\.',oldcover,re.S)[0]
    cover=head+r'''\begin{document}\sloppy
29 October 2026\\
Editor-in-Chief\\
\emph{Human Reproduction}

Dear Editor-in-Chief,

Please consider our manuscript, ``'''+fulltitle+r'''", as an Original Article.

We ask how changing country population shares within age contribute to a changing China-minus-international-reference difference in modeled PCOS-attributed infertility, even when age-standardization weights remain fixed. This informs the interpretation of pooled international benchmarks used in reproductive-health monitoring.

In GBD 2023 estimates for females aged 15--49, the total China-minus-18-country difference changed by +151.54 per 100,000 between 1990 and 2023. Established symmetric accounting allocated +88.65 to reference-country composition, +271.51 to Chinese rates and $-208.62$ to reference-country rates. Substituting UN WPP 2024 shares gave +86.76 for composition. Alternative benchmark definitions changed the sign of a near-zero secondary-infertility contrast in 2023.

Liu and colleagues' 2024 \emph{Human Reproduction} study (doi:10.1093/humrep/dead241) examined burden and geographical trends. Our contribution addresses pooled-reference construction and reporting, rather than a new decomposition algorithm, clinical association or allegation that earlier national comparisons were erroneous. This exploratory aggregate analysis does not estimate causal effects, individual risk or treatment benefit. Confirmed jointly aligned GBD posterior draws were unavailable; derived uncertainty intervals are therefore not supplied. Code and aggregate results are available at an immutable public revision identified in the article; licensed GBD originals and restricted source-cell extracts are excluded.

'''+AI+'\n\n'+declaration+r'''

Yours sincerely,\\
XX Corresponding Author\\
XX Institution; XX email; ORCID: XX
\end{document}
'''
    cover=cover.replace(fulltitle+r'", as',fulltitle+"'', as")
    write(NEW/'manuscript/cover_letter.tex',cover)
    for name in ['main','supplement','cover_letter']:
        before=(OLD/'manuscript'/f'{name}.tex').read_text();after=(NEW/'manuscript'/f'{name}.tex').read_text()
        assert before.count('XX')==after.count('XX'),(name,'XX count changed')
    evidence=json.loads((PRIVATE/'session_model_evidence.json').read_text())
    assert evidence['records'][-1]['model']=='gpt-6.1-sol'
    save(NEW/'docs/AI_VERSION_EVIDENCE.json',{'verified_tool':'OpenAI Codex','recorded_model_identifier':'gpt-6.1-sol','verified_date':'2026-10-29','basis':'turn_context metadata of the current PCOS editing session','evidence_sha256':sha(PRIVATE/'session_model_evidence.json'),'earlier_separate_ChatGPT_versions':'not independently verified; not supplied','exact_serving_snapshot':'not established by the model identifier','author_final_verification':'pending; no approval assumed'})
    save(NEW/'docs/DATA_CITATION_TARGETS.json',{'immutable_research_revision':SNAPSHOT,'code_results_url':STABLE,'un_source_url':UN_URL,'un_source_manifest':json.loads((ROOT/'PCOS_Revision_2026-10-29/data/input/UN_WPP2024_source_manifest.json').read_text()),'restricted_originals_public':False})
    write(NEW/'README.txt','Human Reproduction v5 — 29 October 2026\n\nIndependent technical revision of v4; no new estimates, data or submission.\nAuthoritative editable TeX: manuscript/. Complete editable Word: word/. TeX-generated reading PDFs: output/. Separate editable tables and TIFF figures: submission/.\nThe four authorized corrections are documented in docs/CHANGES_AND_REMAINING_ZH.txt. XX author declarations remain pending.\nAudited research code/results are pinned to '+STABLE+'.\nBuild entry point: scripts/build_hr_v5_20261029.py. The preserved v4 Word converter is reused without modifying it.\n')
    write(NEW/'docs/CHANGES_AND_REMAINING_ZH.txt','Human Reproduction v5 技术修改记录\n\n1. 主稿TeX与Word标题页显式加入42字符Running title。\n2. AI使用声明同时置于主稿致谢与附信，Methods保留简述；仅当前PCOS Codex会话模型标识gpt-6.1-sol经元数据核实。此前独立ChatGPT会话版本未知，不猜测。作者最终人工核查、责任及批准仍待确认。\n3. GBD/UN/代码与派生结果以[dataset]格式标注，增加UN具体文件及不可变研究快照。快照是已核验的分析版本，不冒充v5稿件自身的提交号；许可GBD及源单元继续不公开。\n4. 补充稿删除协作搜索/修订历史/按用户要求等制作措辞，并用全题标识对应主稿。\n\n保留v4全部文件及所有既定XX。原有数值、7个编号公式、13表和12图不变；未重跑统计。Word/PDF将完整渲染核查。\n\n仍待作者确认：姓名顺序、单位及ORCID、贡献/签字/最终批准、经费与披露、伦理适用性、独家审稿与相关稿件、AI人工核查和未知历史版本资料。正式投稿时再验证实时上传格式及系统合并PDF，避免重复表图；本轮不登录或上传投稿系统。\n')

def converter():
    c=module('v4_word_converter',ROOT/'scripts/convert_hr_v4_word_20261029.py')
    c.SOURCE=NEW/'manuscript';c.OUT=NEW/'word';c.PRIVATE=PRIVATE/'word';c.MAIN='PCOS_HumanReproduction_v5_Manuscript.docx';c.SUPP='PCOS_HumanReproduction_v5_Supplement.docx'
    oldprepare=c.prepare
    def prepare(kind):
        ast,info,s=oldprepare(kind)
        if kind=='main':
            matches=[b for b in ast['blocks'] if c.plain(b).startswith('Running title:')]
            assert len(matches)==1
            ast['blocks'].remove(matches[0]);ast['blocks'].insert(1,matches[0])
        return ast,info,s
    c.prepare=prepare
    return c

def documents():
    marker('docx',4)
    c=converter();c.build()
    b=module('hr_v4_builder',ROOT/'scripts/build_hr_v4_20261029.py');b.NEW=NEW
    b.make_docs((NEW/'manuscript/main.tex').read_text(),(NEW/'manuscript/cover_letter.tex').read_text(),False)
    clean_tables()
    from docx import Document
    from docx.shared import Mm
    d=Document(NEW/'submission/Cover_Letter.docx');d.sections[0].page_width=Mm(210);d.sections[0].page_height=Mm(297);d.save(NEW/'submission/Cover_Letter.docx')

def clean_tables():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    d=Document(NEW/'submission/Tables.docx')
    for st in d.styles:
        if st.type==1:
            st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
            for node in st.element.xpath('.//w:rFonts'):
                for attr in list(node.attrib):
                    if attr.endswith('Theme'):del node.attrib[attr]
            for node in st.element.xpath('.//w:pBdr'):node.getparent().remove(node)
    for node in d.element.xpath('.//w:pBdr'):node.getparent().remove(node)
    d.save(NEW/'submission/Tables.docx')

def compile_tex():
    marker('pdf',3)
    out=NEW/'output';out.mkdir(exist_ok=True)
    for source,name in [('main','PCOS_HumanReproduction_v5'),('supplement','PCOS_HumanReproduction_v5_Supplement'),('cover_letter','HumanReproduction_Cover_Letter_v5')]:
        subprocess.run(['/opt/homebrew/bin/tectonic','--keep-logs','--outdir',str(out),str(NEW/'manuscript'/f'{source}.tex')],cwd=NEW/'manuscript',check=True)
        shutil.move(out/f'{source}.pdf',out/f'{name}.pdf')

def verify():
    c=converter();c.verify()
    baseline=json.loads((NEW/'docs/V4_PRESERVATION_BASELINE.json').read_text());assert all(sha(ROOT/p)==h for p,h in baseline.items())
    maps={k:json.loads((NEW/'word'/f'{k}_content_map.json').read_text()) for k in ['main','supplement']}
    for kind in maps:
        old=json.loads((OLD/'word_2026-10-29'/f'{kind}_content_map.json').read_text());new=maps[kind]
        assert old['all_math_latex']==new['all_math_latex']
        assert [t['rows'] for t in old['tables']]==[t['rows'] for t in new['tables']]
        assert [f['sha256'] for f in old['figures']]==[f['sha256'] for f in new['figures']]
        assert old['xx_count']==new['xx_count']
    from docx import Document
    d=Document(NEW/'word'/c.MAIN);text='\n'.join(p.text for p in d.paragraphs)
    assert 'Running title: '+RUNNING in text and len(RUNNING)==42
    assert text.count('[dataset]')==3
    cover=Document(NEW/'submission/Cover_Letter.docx');covertext='\n'.join(p.text for p in cover.paragraphs)
    assert 'gpt-6.1-sol' in covertext
    assert covertext.count('XX')==(OLD/'manuscript/cover_letter.tex').read_text().count('XX')
    assert not any(s in (NEW/'manuscript/supplement.tex').read_text() for s in ["at the user's request",'connected Drive search','This revision changes scientific framing'])
    save(NEW/'docs/TECHNICAL_VERIFICATION.json',{'verified_at':datetime.now().astimezone().isoformat(),'v4_files_preserved':len(baseline),'all_original_XX_preserved':True,'all_original_formulas_tables_figures_unchanged':True,'running_title_characters':42,'dataset_citations':3,'code_results_snapshot':SNAPSHOT,'ai_version_scope':'current documented Codex session only','author_declarations_pending':True,'statistics_rerun':False,'visual_inspection':'pending','submitted':False})

def package():
    paths=[p for p in NEW.rglob('*') if p.is_file() and p.suffix not in ['.zip','.log'] and p.name!='DELIVERY_MANIFEST.json']
    paths += [Path(__file__).resolve(),ROOT/'scripts/convert_hr_v4_word_20261029.py',ROOT/'scripts/build_hr_v4_20261029.py']
    save(NEW/'docs/DELIVERY_MANIFEST.json',{'files':[{'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':sha(p),'git_blob':subprocess.check_output(['git','hash-object',str(p)],text=True).strip()} for p in sorted(set(paths))],'restricted_raw_data_included':False})
    paths.append(NEW/'docs/DELIVERY_MANIFEST.json')
    with zipfile.ZipFile(NEW/'output/PCOS_HumanReproduction_v5_Complete.zip','w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(set(paths)):z.write(p,p.relative_to(ROOT))
    print('Archive members:',len(set(paths)))

def finalize():
    verify()
    qa=json.loads((NEW/'docs/VISUAL_QA.json').read_text())
    assert qa['total_pages_inspected']==84
    assert all(d['all_pages_visually_inspected'] for d in qa['documents'].values())
    for p in [NEW/'docs/TECHNICAL_VERIFICATION.json',NEW/'word/CONVERSION_QC.json']:
        d=json.loads(p.read_text());d['visual_inspection']='passed all 84 pages (45 Word-render pages and 39 TeX PDF pages)';save(p,d)
    d=json.loads((NEW/'docs/TECHNICAL_VERIFICATION.json').read_text())
    d['word_pages']={'main':21,'supplement':20,'cover':1,'tables':3}
    d['tex_pdf_pages']={'main':19,'supplement':19,'cover':1}
    save(NEW/'docs/FINAL_VERIFICATION.json',d)
    package()

if __name__=='__main__':
    action=sys.argv[1]
    if action=='build':source();documents();compile_tex();verify()
    else:{'documents':documents,'tables-style':clean_tables,'compile':compile_tex,'verify':verify,'package':package,'finalize':finalize}[action]()
