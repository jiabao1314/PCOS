#!/usr/bin/env python3
"""Independent HR v6 editorial revision from preserved v5 and verified aggregates.

Run with the bundled document Python. No upstream GBD fit, synthetic draws,
patient records, Git mutation, cloud write or journal submission occurs here.
"""
import argparse
import csv
import difflib
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import types
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
OLD = ROOT / 'PCOS_HumanReproduction_v5_2026-10-29'
NEW = ROOT / 'PCOS_HumanReproduction_v6_2026-10-30'
PRIVATE = ROOT / '.shared_context_runtime/v6'
PRE = ROOT / 'PCOS_Revision_Preflight_2026-10-30'
SCI = ROOT / 'PCOS_Revision_2026-10-29'
P = '20261030_'
BASE = P + 'PCOS_HumanReproduction_v6_'
TEX = {k: NEW / 'manuscript' / (BASE + n + '.tex') for k, n in [('main','Manuscript'),('supplement','Supplement'),('cover','CoverLetter')]}
WORD = {k: NEW / 'word' / (BASE + n + '.docx') for k, n in [('main','Manuscript'),('supplement','Supplement'),('cover','CoverLetter'),('response','CommentResponse'),('tables','Tables')]}
PDF = {k: NEW / 'output' / (BASE + n + '.pdf') for k, n in [('main','Manuscript'),('supplement','Supplement'),('cover','CoverLetter'),('response','CommentResponse')]}
SNAPSHOT = '1abf337d9f23cc69292e03cd581b6a1557e6effe'
EXPECTED_IDS = ['0','3','4','5','8','15','18','19','22','25','26','31','36','39','48','51','76']

ABSTRACT = r'''\noindent\textbf{STUDY QUESTION:} How did changing reference-country population shares contribute to China's PCOS-attributed infertility rate difference from an 18-country pool during 1990--2023?

\noindent\textbf{SUMMARY ANSWER:} Reference-country composition contributed +88.65 per 100,000 females to the +151.54 change in the modeled difference, or 58.5\% of its signed net change.

\noindent\textbf{WHAT IS KNOWN ALREADY:} GBD studies describe PCOS-attributed infertility burden and trends. Fixed age-standardization weights still permit country composition within each reference age group to change.

\noindent\textbf{STUDY DESIGN, SIZE, DURATION:} This exploratory ecological analysis covered China and 18 other sovereign G20 countries, females aged 15--49 in seven age groups, and 34 years, 1990--2023.

\noindent\textbf{PARTICIPANTS/MATERIALS, SETTING, METHODS:} GBD 2023 prevalence counts and matched female populations yielded age-specific rates. Fixed age weights were normalized over ages 15--49. The primary estimand was the change in China's total-infertility standardized rate difference from the contemporaneous population-pooled reference. Symmetric accounting separated Chinese rates, reference-country rates and reference-country shares. UN WPP 2024 replaced shares only; subtypes and alternative benchmarks addressed secondary questions. Sensitivity analyses covered country exclusions, age restriction, alternative endpoints, annual chaining and all 128 rounding-perturbation corners of published age weights.

\noindent\textbf{MAIN RESULTS AND THE ROLE OF CHANCE:} China's total rate rose from 333.56 to 605.08 and the reference rate from 606.48 to 726.45 per 100,000. Their difference narrowed from $-272.92$ to $-121.38$. The respective rate and composition components were $+271.51-208.62+88.65=+151.54$ (rounded). UN shares gave +86.76 for composition and +150.69 overall. The 2023 secondary-infertility difference was +2.16 with contemporaneous GBD shares, $-47.99$ with fixed-1990 shares and $-127.89$ with equal-country weights. Results are derived point estimates; jointly aligned posterior draws were not obtained for interval propagation.

\noindent\textbf{LIMITATIONS, REASONS FOR CAUTION:} Uncertainty in component magnitudes and near-zero contrast signs remains unquantified. The signed net denominator contains opposing components, so 58.5\% is an accounting ratio. Rounded published age weights and the economically selected G20 reference also limit interpretation.

\noindent\textbf{WIDER IMPLICATIONS OF THE FINDINGS:} Reproductive-health monitoring against pooled populations should report reference weights and absolute components alongside net differences. Doing so distinguishes changes in national rates from changes in the population represented by the benchmark.

'''

INTRO = r'''\section{Introduction}
Polycystic ovary syndrome (PCOS) affects reproductive health and requires care tailored to the individual \citep{pcos}. Population monitoring complements this clinical perspective by describing national trajectories in PCOS-attributed infertility. In the Global Burden of Disease (GBD) study, this outcome is a modeled cause-specific infertility impairment expressed in the general female population \citep{gbd}. Interpreting a national trajectory against a pooled international population requires an explicit definition of that reference.

Liu and colleagues quantified GBD 2019 PCOS-attributed infertility prevalence, disability and geographical trends \citep{liu_pcos2024}. Liang and colleagues analysed GBD 2021 burden, trends and the contributions of population growth, ageing and epidemiological change to burden counts \citep{liang_burden2025}; another study examined China and worldwide trajectories \citep{liang_trend2026}. These studies establish the burden context. Our question concerns the construction of a particular national-versus-pool rate contrast, with country shares changing within age even when the age standard remains fixed (Supplementary Table~S9).

A contemporaneous population pool represents the country mix at each date. If countries with lower age-specific rates gain population share, the pooled rate can fall even with unchanged country-specific rates. This composition term operates within age and therefore persists after age standardization. We quantified the 1990--2023 change in China's total PCOS-attributed infertility rate difference from the population-pooled reference of 18 other sovereign G20 countries, allocating that change to Chinese rates, reference-country rates and reference-country shares. Secondary questions concern demographic-source choice, infertility subtypes and fixed-1990 or equal-country benchmarks.

'''

METHODS = r'''\section{Materials and methods}
\subsection{Design data and outcome}
This ecological secondary analysis used GBD 2023 Results Tool exports obtained on 25 October 2026 \citep{results}. We analysed females aged 15--49 in seven five-year age groups during 1990--2023. The fixed set comprised 19 sovereign G20 countries, with China excluded from the reference. The 18 reference countries were Argentina, Australia, Brazil, Canada, France, Germany, India, Indonesia, Italy, Japan, Republic of Korea, Mexico, Russian Federation, Saudi Arabia, South Africa, T\"urkiye, United Kingdom and United States of America. The G20 was established in 1999; we applied this country set retrospectively across the study years. The European Union and African Union were excluded as additional units to avoid overlapping country populations \citep{g20}.

We selected prevalence (measure 5), PCOS (cause 605), total infertility (impairment 198), primary infertility (220) and secondary infertility (221). Endometriosis (607) was retained as a supplementary comparator. The GBD 2023 household-survey framework identifies primary infertility using absence of a livebirth, desire for a child, union duration of at least five years and no current contraception. Its secondary definition requires a livebirth more than five years earlier and the corresponding union, fertility-intention and contraception conditions. The modeling process combines these survey envelopes with sex- and cause-attribution evidence, including clinical sources, and expresses the final rates among all females in each age group \citep{gbd}. Supplementary data distinguish these quantities from clinical 12-month infertility definitions and from PCOS prevalence itself.

For country $c$, age $a$, year $t$ and infertility type $k$, let $N_{c,a,t,k}$ denote modeled infertility cases and $P_{c,a,t}$ the matched GBD female population. Age-specific prevalence per 100,000 is $r_{c,a,t,k}=10^5N_{c,a,t,k}/P_{c,a,t}$. We checked unique complete cells, finite nonnegative values, primary-plus-secondary additivity and agreement with separately exported endpoint Rate values. Input specifications and hashes are retained in the reproducibility archive.

\subsection{Annual contrasts and the primary period change}
Let $\mathcal{R}$ be the set of 18 reference countries; the labels $\mathrm{CHN}$ and $\mathrm{R}$ denote China and the pooled reference. We normalized the seven published age weights over ages 15--49 and held them fixed. Their six-decimal reproduction is attributed to Chen and colleagues \citep{chen}; an official full-precision IHME vector was not independently verified.
\begin{align}
 s_{c,a,t}&=\frac{P_{c,a,t}}{\sum_{j\in\mathcal{R}}P_{j,a,t}}, &
 A_{\mathrm{R},t,k}&=\sum_a w_a\sum_{c\in\mathcal{R}}s_{c,a,t}r_{c,a,t,k},\\
 A_{\mathrm{CHN},t,k}&=\sum_a w_ar_{\mathrm{CHN},a,t,k}, &
 D_{t,k}&=A_{\mathrm{CHN},t,k}-A_{\mathrm{R},t,k}.
\end{align}
Here $w_a$ sums to one, and $s_{c,a,t}$ sums to one over reference countries within each age and year. Pooling within age precedes standardization. The annual rate difference is $D_{t,k}$; its period change is $\Delta D_k(t_0,t_1)=D_{t_1,k}-D_{t_0,k}$. The primary estimand is $\Delta D_{\mathrm{total}}(1990,2023)$ under contemporaneous GBD shares. Subtype and alternative-reference analyses address secondary estimands. All rates are restricted-age standardized rates per 100,000 females.

\subsection{Symmetric decomposition of the period change}
For endpoints $t_0$ and $t_1$, the symmetric product identity of Kitagawa-type accounting gives \citep{kitagawa}:
\begin{align}
 \Delta D_k(t_0,t_1)&=X_k+Q_k+S_k,\\
 X_k&=\sum_a w_a(r_{\mathrm{CHN},a,t_1,k}-r_{\mathrm{CHN},a,t_0,k}),\\
 Q_k&=-\sum_a w_a\sum_{c\in\mathcal{R}}\frac{s_{c,a,t_0}+s_{c,a,t_1}}{2}(r_{c,a,t_1,k}-r_{c,a,t_0,k}),\\
 S_k&=-\sum_a w_a\sum_{c\in\mathcal{R}}\frac{r_{c,a,t_0,k}+r_{c,a,t_1,k}}{2}(s_{c,a,t_1}-s_{c,a,t_0}).
\end{align}
$X_k$, $Q_k$ and $S_k$ allocate the period change to Chinese rates, reference-country rates and reference-country shares, respectively. Each uses the same endpoints, subtype and population-source specification. Averaging the two replacement orders splits the product interaction symmetrically, giving an exact identity before rounding. Positive components raise the signed China-minus-reference difference. When the difference is negative, an upward change narrows it.

The identity is additive across ages and types. The share component is zero if shares are unchanged; it also vanishes when the countries' mean endpoint rates are identical within age, because share changes sum to zero. These algebraic properties clarify the role of composition. Its estimated magnitude in the GBD data is reported separately below.

We report absolute components before the ratio $100S_k/\Delta D_k(t_0,t_1)$. Its signed net denominator contains cancellation between components and can approach zero; ratios can consequently exceed 100\% or change sign. This accounting ratio is distinct from a fraction of women, cases or disease causation. Subtype and country-share contributions partition the same change in different ways and are not additive explanations.

\subsection{Population sources and benchmark definitions}
UN WPP 2024 supplied historical mid-year female population estimates for the same 4,522 country--age--year cells \citep{wpp2024}. Values in thousands were converted to persons; projections from 2024 onward were excluded. Only the within-age reference shares were replaced. GBD cases continued to use matched GBD population denominators, and the age standard remained fixed.

Current-year shares compare China with the contemporaneous pooled population. Fixed-1990 shares hold the reference's 1990 within-age country mix constant while country rates vary. Equal-country weights compare China with the arithmetic mean of 18 country rates within age. These benchmarks specify different reference populations and therefore different descriptive estimands. A fixed-share path is distinct from the symmetric share component.

\subsection{Structural sensitivity and posterior uncertainty}
Existing sensitivity analyses used both population sources, each country exclusion, ages 20--49 with renormalized weights, 1990 or 2010 start years, and 2021, 2022 or 2023 endpoints. Adjacent-year decompositions were summed over 1990--2023. All $2^7=128$ perturbation corners of the unnormalized age weights by $\pm0.5\times10^{-6}$ assessed rounding. These are scenario and numerical-precision checks. Supplementary endometriosis analyses include all annual subtype differences and recent endpoint comparisons, supporting descriptive examination of cancellation \citep{xu}.

The available study exports contain marginal summaries without posterior draw columns. The documented access audit did not identify a provider-confirmed joint array aligned across countries, ages, years and infertility types. Derived results therefore remain point estimates without propagated uncertainty intervals or significance tests. Marginal bounds, independent simulation and country bootstrap do not recover the required joint dependence. Eligible draws would require provider alignment evidence and product-summary checks before recomputing pooling and decomposition within each draw. Ratios would additionally require diagnostics for a near-zero net-change denominator.

\subsection{Computational reproducibility}
The exploratory analyses were not externally preregistered. Core results were independently calculated in Python 3.12.14 and base R 4.6.1. The recorded Python scientific environment used NumPy 2.3.5, pandas 2.2.3 and Matplotlib 3.10.7. These are external libraries, separate from the language's standard library. The source manifest, code and aggregate results are pinned to the immutable research snapshot \citep{repository}. The supplement provides the code-variable notation map, commands, numerical checks and GATHER reporting map \citep{gather}.

'''

DISCUSSION = r'''\section{Discussion}
Reference-country composition made a substantial contribution to the narrowing PCOS-attributed infertility contrast between China and the pooled reference. With a fixed age standard, the 1990--2023 signed rate difference increased by +151.54 per 100,000: +271.51 from Chinese rates, $-208.62$ from reference-country rates and +88.65 from changing reference-country shares. The close +86.76 share component under UN populations supports stability to this demographic-source substitution. Both China's and the reference's total modeled rates rose, so the narrower contrast describes their relative trajectories alongside the changing population mix.

This analysis adds a reference-construction perspective to geographical burden and trend studies \citep{liu_pcos2024,liang_burden2025,liang_trend2026}. Holding age weights fixed retains variation in country shares within age. Explicit decomposition makes that source of change visible in the national-versus-pool comparison. Established symmetric accounting gives the absolute contributions in the same rate units as the observed difference. This application concerns an explicitly defined pool and period; the prior studies address the targets summarized in Supplementary Table~S9.

The practical implication is to specify the population represented by a benchmark and show its weights with the comparison. Contemporary pooling addresses the country mix at each date, fixed historical shares address a constant mix, and equal-country weighting gives small and large countries equal influence within age. The 2023 secondary-infertility point estimate is positive under contemporary shares and negative under the alternatives. Reporting these targets together reveals the benchmark dependence of that near-zero contrast. Absolute components should accompany ratios: the 58.5\% share ratio uses a net denominator reduced by opposing rate components, and its 39.8\%--67.9\% range across the specified age and country-exclusion scenarios describes sensitivity to those choices.

Clinical interpretation also depends on the outcome and denominator. Yang and colleagues' Chinese community survey used Rotterdam PCOS ascertainment among married Han women aged 20--49 and reported infertility characteristics among PCOS participants \citep{yang}. The present analysis instead concerns GBD cause-attributed infertility among all females aged 15--49. These definitions and populations serve different purposes (Supplementary Table~S11). The clinical guideline provides the diagnostic and care context \citep{pcos}; the aggregate comparison does not identify individual transitions between infertility categories or effects of diagnosis, treatment or service provision.

The principal limitation is unquantified joint posterior uncertainty. Without verified aligned draws, the precision of absolute components, the 58.5\% ratio and near-zero contrast signs remains unknown; scenario analyses cannot supply that precision. Additional limits include heterogeneous upstream evidence and modeled attribution, possible shared population inputs between GBD and UN, rounded published age weights, exploratory analytical choices and the economically selected country set. The results describe this G20 comparison rather than the world or a matched clinical control. The computational audits check implementation and source-output consistency, while clinical validity requires evidence matched to the modeled definitions.

Reproductive-health monitoring can be made more interpretable by reporting the signed national-versus-reference change together with absolute rate and country-composition components. In these estimates, the reference population's changing composition accounts for +88.65 per 100,000 of a +151.54 change, providing information that the net contrast alone would conceal.

'''

SUPP_ADDITIONS = r'''
\clearpage
\section*{Annual subtype differences and recent endpoints}
The machine-readable annual table contains 408 rows: two causes, two population sources, three benchmark definitions and 34 years. Within every cause--source--benchmark--year combination, the total difference equals the sum of primary and secondary differences; maximum absolute residual is $6.82\times10^{-13}$ per 100,000. Table~\ref{tab:s10} displays the recent endometriosis endpoints under contemporary shares. The GBD 2023 total contrast of $-0.202454$ combines primary $-7.935717$ and secondary $+7.733263$. In 2022 both subtype differences were negative. This locates the endpoint cancellation in the actual annual output. The annual model point estimates are temporally related; no ordinary correlation test treating 34 years as independent observations was used.

\small
\begin{longtable}{rlrrr}
\caption{Endometriosis China-minus-reference differences by infertility type at recent endpoints.}\label{tab:s10}\\
\toprule Year & Shares & Primary & Secondary & Total\\\midrule\endhead
\input{tables/20261030_v6_s10_endometriosis_endpoints.tex}
\end{longtable}\normalsize
All differences are per 100,000 females with identical age weights. UN substitutes only within-age country shares. Values are model point estimates and sum before displayed rounding. Complete 1990--2023 annual results accompany the article as a CSV.

\section*{Clinical and modeled outcome definitions}
The GBD 2023 methodological appendix distinguishes household-survey definitions from clinic-based evidence and describes exposure, sex attribution and cause allocation. Primary household-survey infertility uses no livebirth, desire for a child, at least five years in union and no current contraception; secondary infertility uses a last livebirth more than five years earlier with the corresponding other conditions. The final sex-specific impairment rates are expressed in the general reproductive-age population. The PCOS model uses an NIH reference definition and crosswalks alternative ascertainment definitions. Infertility-related disease sequelae are then allocated to primary and secondary envelopes. Exact sources are Appendix 1 printed pp. 590--591 and 684--692 (PDF pp. 591--592 and 685--693).

\small
\begin{longtable}{p{.20\textwidth}p{.36\textwidth}p{.34\textwidth}}
\caption{Outcome and population distinctions for the interpretation of PCOS-related evidence.}\label{tab:s11}\\
\toprule Evidence & Outcome and denominator & Use in this study\\\midrule\endhead
GBD 2023 impairment model & Cause-attributed primary/secondary infertility; all females aged 15--49 in each modeled age group. & Supplies the national and pooled-reference rates. Household-survey envelopes and clinical inputs have distinct definitions.\\
GBD 2023 PCOS model & NIH reference ascertainment; alternative definitions crosswalked before disease and sequela estimation. & Establishes upstream model context; PCOS prevalence itself is a different outcome.\\
Yang et al., 2022 & Rotterdam PCOS ascertainment in married Han women aged 20--49; infertility characteristics reported within PCOS participants. & Provides clinical/population context. The detailed infertility time-at-risk rule was not independently established from a questionnaire for this revision.\\
Teede et al., 2023 guideline & Clinical PCOS assessment and management; diagnostic recommendations are distinct from modeled infertility impairment prevalence. & Supports diagnostic and care context rather than a disease-rate validation dataset.\\\bottomrule
\end{longtable}\normalsize
Bibliographic details and DOIs are in the main reference list. No clinical microdata were obtained or harmonized with the GBD estimates. The Yang survey is used for definition comparison, without a numerical validation claim.

\section*{Notation and software mapping}
$N_{c,a,t,k}$ denotes infertility cases and maps to \texttt{gbd\_cases}; $P_{c,a,t}$ maps to \texttt{gbd\_pop}. China is labeled $\mathrm{CHN}$, the reference set is $\mathcal{R}$, and $\mathrm{R}$ labels its aggregate. The code's \texttt{gbd\_rate} is $r_{c,a,t,k}$; \texttt{china\_aspr}, \texttt{ref\_aspr} and \texttt{gap} are $A_{\mathrm{CHN},t,k}$, $A_{\mathrm{R},t,k}$ and $D_{t,k}$. \texttt{net\_gap\_change} is $\Delta D_k(t_0,t_1)$; \texttt{china\_rate\_effect}, \texttt{reference\_rate\_effect} and \texttt{reference\_country\_composition\_effect} are $X_k$, $Q_k$ and $S_k$. All terms in a decomposition row share endpoints and population-source weights.

The archived scientific environment records Python 3.12.14, NumPy 2.3.5, pandas 2.2.3, Matplotlib 3.10.7 and base R 4.6.1 on aarch64 macOS. The document-generation environment is recorded separately in the dated verification files. From the scientific snapshot directory, the executable sequence is \texttt{python code/PCOS\_NewArticle\_complete.py}, \texttt{Rscript code/PCOS\_NewArticle\_complete.R} and \texttt{python code/revision\_analysis.py}. These commands write results and should be run in a separate working copy with verified inputs. This revision reused the validated results; it did not repeat the upstream GBD fit. Core/sensitivity computations, subsequent numerical checks and this editorial revision are dated in the archived run records; the analyses were exploratory and not preregistered.
'''


def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def save(p, value):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
def write(p, text):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text)
def between(text, start, end, replacement):
    a, b = text.index(start), text.index(end)
    return text[:a] + replacement + text[b:]
def module(name, path):
    spec=importlib.util.spec_from_file_location(name,path)
    obj=importlib.util.module_from_spec(spec);spec.loader.exec_module(obj);return obj


def sources():
    assert not NEW.exists(), 'Preserve any existing revision'
    NEW.mkdir()
    for d in ['manuscript','word','output','docs','results','submission']:
        (NEW/d).mkdir()
    baseline={str(p.relative_to(ROOT)):sha(p) for p in OLD.rglob('*') if p.is_file()}
    save(NEW/'docs'/ (P+'v5_preservation_baseline.json'),baseline)
    # New asset filenames are dated; only unchanged bytes are copied.
    asset_map={}
    for section in ['tables','figures']:
        for p in (OLD/'manuscript'/section).iterdir():
            if not p.is_file():continue
            target=NEW/'manuscript'/section/(P+'v6_'+p.name)
            target.parent.mkdir(exist_ok=True)
            shutil.copy2(p,target)
            asset_map[section+'/'+p.name]=section+'/'+target.name
    for p in (OLD/'submission/figures').glob('*.tif'):
        dest=NEW/'submission/figures'/(P+'v6_'+p.name);dest.parent.mkdir(exist_ok=True);shutil.copy2(p,dest)
    main=(OLD/'manuscript/main.tex').read_text().replace('HUMAN_REPRODUCTION_V5_2026_10_29','HUMAN_REPRODUCTION_V6_2026_10_30').replace('29 October 2026','30 October 2026',1)
    a=main.index(r'\noindent\textbf{STUDY QUESTION:}')
    b=main.index(r'\noindent\textbf{FUNDING:}',a)
    main=main[:a]+ABSTRACT+main[b:]
    main=between(main,r'\section{Introduction}',r'\section{Materials and methods}',INTRO)
    main=between(main,r'\section{Materials and methods}',r'\subsection{Ethical approval}',METHODS)
    main=between(main,r'\subsection{Use of artificial intelligence}',r'\section{Results}',r'''\subsection{Use of artificial intelligence}
OpenAI ChatGPT and Codex supported code development, analysis checks, drafting and technical editing, including this revision on 30 October 2026. Acknowledgements records the verified earlier model identifier, version limits and pending author-review status.

''')
    result_start=main.index(r'\section{Results}')
    result_end=main.index(r'\section{Discussion}',result_start)
    result=main[result_start:result_end]
    start=result.index('The GBD population and UN subset')
    end=result.index("China's PCOS-attributed",start)
    result=result[:start]+'''The GBD population and UN subset each covered 4,522 cells. The prevalence analysis contained 27,132 country--age--year--cause--type cells across the two causes, including 13,566 PCOS cells. Completeness, subtype additivity, endpoint rate reconstruction and independent R/Python checks passed; numerical details are in Supplementary data.

'''+result[end:]
    result=result.replace('Both rates increased, so a narrowing negative contrast does not establish improved reproductive health.','The narrower negative difference therefore accompanied increases in both standardized rates.')
    result=result.replace('A substantial contribution persisted, but it was not a majority in every specification. These limits describe specified benchmark sensitivities and are not an uncertainty interval.','The contribution persisted across these specifications, with its ratio to the net change varying with the age range and reference countries.')
    result=result.replace('The 69.4\\% subtype ratio and 58.5\\% country-share ratio describe intersecting partitions of the same change and must not be added.','These ratios describe intersecting partitions of the same total change, as defined in Methods.')
    result=result.replace('Similar results under this alternative demographic source support stability to population-source choice, not independent accuracy of GBD infertility rates.','The share and net-change components were close under the two population sources.')
    result=result.replace('Thus the small positive secondary contrast does not support a stable sign or international ranking. The three benchmarks describe distinct reference populations, rather than competing estimates of one clinical risk.','The 2023 secondary point estimate was positive under contemporary shares and negative under both alternatives, demonstrating the dependence of this contrast on the reference definition.')
    result=result.replace('percentages are not disease fractions or causal effects.','percentages use the signed net denominator defined in Methods.')
    result=result.replace('Scenario differences are not posterior intervals.','Each row represents a different benchmark definition.')
    main=main[:result_start]+result+main[result_end:]
    main=between(main,r'\section{Discussion}',r"\section*{Authors' roles}",DISCUSSION)
    old_ai=re.search(r'OpenAI ChatGPT and Codex supported code development, analysis checks, drafting and technical editing under researcher-specified.*?human authors retain responsibility for the work\.',main,re.S)[0]
    ai=r'''OpenAI ChatGPT and Codex supported code development, analysis checks, drafting and technical editing under researcher-specified questions and revision instructions, including editorial revision and document checks on 30 October 2026. The earlier documented Codex session on 29 October 2026 recorded \texttt{gpt-6.1-sol}; exact model versions for earlier separate ChatGPT sessions and an independently verified serving identifier for this later revision are not specified. No artificial observations or posterior draws replaced study data. The authors' final scientific review, verification of AI-supported content and approval remain to be confirmed before submission; human authors retain responsibility for the work.'''
    main=main.replace(old_ai,ai)
    main=main.replace('Supplementary methods, eight numerical tables, seven figures, a literature/estimand comparison table (S9), and a GATHER reporting map accompany the article.','Supplementary methods, nine numerical tables (S1--S8 and S10), seven figures, the literature/estimand comparison (S9), the outcome-definition comparison (S11), and a GATHER reporting map accompany the article. A complete annual subtype CSV and dated source, notation and revision audits are included in the reproducibility archive.')
    bibliography=r'''\bibitem[Australian Government(2026)]{g20}Australian Government Department of Foreign Affairs and Trade. The G20. Accessed 30 October 2026. \url{https://www.dfat.gov.au/trade/organisations/g20}.
'''
    main=main.replace(r'\begin{thebibliography}{99}',r'\begin{thebibliography}{99}'+'\n'+bibliography)
    yang=r'''\bibitem[Yang et al.(2022)]{yang}Yang R, Li Q, Zhou Z, et al. Changes in the prevalence of polycystic ovary syndrome in China over the past decade. \emph{Lancet Reg Health West Pac}. 2022;25:100494. \href{https://doi.org/10.1016/j.lanwpc.2022.100494}{doi:10.1016/j.lanwpc.2022.100494}.
'''
    main=main.replace(r'\end{thebibliography}',yang+r'\end{thebibliography}')
    supp=(OLD/'manuscript/supplement.tex').read_text().replace('HUMAN_REPRODUCTION_V5_SUPPLEMENT','HUMAN_REPRODUCTION_V6_SUPPLEMENT').replace('29 October 2026','30 October 2026')
    supp=supp.replace('This targeted comparison supports scientific positioning; it is not an exhaustive systematic review or proof that no prior analysis has used similar accounting. Bibliographic details and source URLs are provided in the main article reference list. No pooled-comparator error is attributed to a paper simply because it reports national age-standardized rates.','This targeted comparison places the present estimand alongside selected prior questions. Bibliographic details and source URLs are in the main reference list. It is a targeted evidence comparison, with access levels indicated below.')
    supp=supp.replace('GBD 2019 global, regional and national prevalence, YLD and trends, 1990--2019. Published Methods describe direct age standardization.','GBD 2019 global, regional and national prevalence, YLD and trends, 1990--2019. Published Methods describe direct age standardization; Methods text and indexed abstract verified.')
    supp=supp.replace('Provides disease-burden context. It does not establish that a China-versus-18-country pool with our within-age weights was used. We do not reinterpret its individual national comparisons.','Provides the geographical disease-burden context. The present analysis specifies a China-versus-18-country pool with explicit within-age country weights.')
    supp=supp.replace('GBD 2021 burden and trends; the Methods describe a Das Gupta allocation to population growth, ageing and epidemiological factors.','GBD 2021 burden and trends; full-text Methods describe a Das Gupta allocation to population growth, ageing and epidemiological factors.')
    supp=supp.replace('Supports a China/worldwide monitoring use case. No claim about an undocumented pooling formula is made; conclusions here do not test that paper\'s estimates.','Supports a China/worldwide monitoring use case. Comparator-construction detail remains unverified at this access level.')
    supp=supp.replace('The GBD input files contain point estimates and marginal lower/upper limits, without posterior draw columns. Availability of a provider-confirmed joint posterior array aligned across all dimensions required by this analysis was not established from the study inputs. This does not imply that suitable draws do not exist or cannot be supplied. The present analyses therefore do not quantify joint posterior uncertainty.','The GBD input files contain point estimates and marginal lower/upper limits without posterior draw columns. The dated access audit on 30 October 2026 did not identify a qualified joint draw product in the study inputs or the bounded accessible-Drive discovery. Official GBD catalogue/results endpoints returned HTTP 403 through the available web route. These observations describe the access attempt, not provider-wide availability. Joint posterior uncertainty remains unquantified.')
    supp=supp.replace('13 & Uncertainty approach & No joint draws; no custom posterior intervals.','13 & Uncertainty approach & Partly met: dated access audit and eligibility criteria; no qualified joint draws or propagated posterior intervals.')
    supp=supp.replace(r'\end{document}',SUPP_ADDITIONS+'\n'+r'\end{document}')
    # Generate the six recent endpoint rows only from the already verified annual CSV.
    import pandas as pd
    annual=pd.read_csv(PRE/(P+'annual_endpoint_check.csv'))
    endpoint=annual.query("cause=='Endometriosis' and benchmark=='dynamic' and year>=2021").sort_values(['year','source'])
    assert len(endpoint)==6
    table=''.join(f"{r.year} & {'GBD' if r.source=='GBD' else 'UN WPP'} & {r.Primary:+.6f} & {r.Secondary:+.6f} & {r.Total:+.6f}"+r'\\'+'\n' for r in endpoint.itertuples())+r'\bottomrule'+'\n'
    write(NEW/'manuscript/tables'/ (P+'v6_s10_endometriosis_endpoints.tex'),table)
    shutil.copy2(PRE/(P+'annual_endpoint_check.csv'),NEW/'results'/(P+'v6_annual_subtype_contrasts.csv'))
    for name in ['source_definition_audit.json','joint_draws_access_audit.json','annual_endpoint_qc.json','country_reference_mapping.csv']:
        shutil.copy2(PRE/(P+name),NEW/'docs'/(P+'v6_'+name))
    # Preserve the original comment source exactly in the response record.
    for old,new in asset_map.items():
        main=main.replace('{'+old+'}','{'+new+'}')
        supp=supp.replace('{'+old+'}','{'+new+'}')
    write(TEX['main'],main);write(TEX['supplement'],supp)
    cover=(OLD/'manuscript/cover_letter.tex').read_text().replace('29 October 2026','30 October 2026',1)
    start=cover.index("Liu and colleagues'")
    end=cover.index('OpenAI ChatGPT',start)
    cover=cover[:start]+r'''The study makes explicit how the population represented by a pooled reference changes within age. It adds absolute country-composition accounting to the geographical burden and trend questions addressed by earlier PCOS-attributed infertility studies, and shows why the same near-zero secondary contrast changes sign under different reference definitions. Reporting reference weights and absolute components can improve the interpretation of reproductive-health monitoring. Jointly aligned posterior draws were not obtained; derived results retain their point-estimate status. Code and aggregate outputs are linked to the immutable scientific snapshot; licensed source extracts remain in the private archive.

'''+cover[end:]
    start=cover.index('OpenAI ChatGPT');end=cover.index('XX: Corresponding author',start)
    cover=cover[:start]+ai+'\n\n'+cover[end:]
    write(TEX['cover'],cover)
    for kind,oldname in [('main','main'),('supplement','supplement'),('cover','cover_letter')]:
        assert TEX[kind].read_text().count('XX')==(OLD/'manuscript'/ (oldname+'.tex')).read_text().count('XX'),kind
    for kind,oldname in [('main','main'),('supplement','supplement'),('cover','cover_letter')]:
        diff=''.join(difflib.unified_diff((OLD/'manuscript'/(oldname+'.tex')).read_text().splitlines(True),TEX[kind].read_text().splitlines(True),fromfile='v5/'+oldname+'.tex',tofile=TEX[kind].name))
        write(NEW/'docs'/(P+'v5_to_v6_'+kind+'.diff'),diff)
    save(NEW/'docs'/(P+'v6_build_provenance.json'),{'created_at':datetime.now(ZoneInfo('Europe/Paris')).isoformat(timespec='seconds'),'scientific_snapshot':SNAPSHOT,'latest_remote_at_start':subprocess.check_output(['git','rev-parse','origin/main'],cwd=ROOT,text=True).strip(),'local_head':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'working_tree_dirty':True,'builder_sha256':sha(Path(__file__)),'converter_sha256':sha(ROOT/'scripts/convert_hr_v4_word_20261029.py'),'point_estimates_changed':False,'qualified_draws_obtained':False,'new_research_data_downloaded':False,'upstream_GBD_refitted':False,'author_XX_preserved':True,'author_final_approval':'pending','submitted':False,'unchanged_assets_renamed_to_dated_v6_names':asset_map})
    responses()


def converter():
    path=ROOT/'scripts/convert_hr_v4_word_20261029.py'
    code=path.read_text().replace("(SOURCE/(kind+'.tex'))","TEX_FILES[kind]").replace("SOURCE/'main.tex'","TEX_FILES['main']").replace("SOURCE/'supplement.tex'","TEX_FILES['supplement']").replace('range(1,10)','range(1,12)')
    needle='        assert len(widths)==n,(idx,n,widths)'
    code=code.replace(needle,"        if info['kind']!='main' and idx==10:widths=[20,32,36,36,36]\n        if info['kind']!='main' and idx==11:widths=[35,60,65]\n"+needle)
    obj=types.ModuleType('v6_converter');obj.__file__=str(path);exec(compile(code,str(path),'exec'),obj.__dict__)
    obj.V4=OLD;obj.SOURCE=NEW/'manuscript';obj.OUT=NEW/'word';obj.PRIVATE=PRIVATE/'word';obj.TEX_FILES=TEX
    obj.MAIN=WORD['main'].name;obj.SUPP=WORD['supplement'].name
    oldprepare=obj.prepare
    def prepare(kind):
        ast, info, source=oldprepare(kind)
        if kind=='main':
            found=[b for b in ast['blocks'] if obj.plain(b).startswith('Running title:')]
            assert len(found)==1
            ast['blocks'].remove(found[0]);ast['blocks'].insert(1,found[0])
        return ast,info,source
    obj.prepare=prepare
    return obj


RESPONSE_ACTIONS = {
'0': ('已验证','重写引言、结果和讨论，以主要发现、绝对分量及监测应用为主线；将必要限制集中到方法和讨论。','主稿 Introduction Results Discussion','结果段不再机械以否定句收尾；主要数值与原结果一致。'),
'3': ('已验证','保留11栏目，缩短主问题、背景和方法；词数按同一空格分词法与v5比较，不把内部编辑目标写成期刊上限。','主稿 Abstract；文档核验词数记录','保留Funding与Disclosures两栏及原XX。'),
'4': ('已验证','明确19个主权国家，其中中国之外18国构成参照；固定集合回溯至1990年，EU/AU不重复纳入。','主稿 Design data and outcome；国家ISO3映射CSV','官方G20说明及4522单元前置审计支持；不临时加入受邀国家。'),
'5': ('已验证','摘要采用derived point estimates，说明未取得对齐draws；压缩重复防御措辞。','主稿 Abstract与Structural sensitivity and posterior uncertainty','按真实无draws分支表述，场景范围不冒充95%UI。'),
'8': ('已验证','保持作者年与研究目标表述，更新S9证据层级和不同估计目标对照。','主稿 Introduction；补充表S9；文献估计目标CSV','无in this journal；Liang2026仅摘要层证据，不推断pooling公式。'),
'15': ('已验证','病例数统一N，中国统一CHN，参照集合使用花体R；年度D与期间ΔD一并统一。','主稿 Annual contrasts及Symmetric decomposition；补充Notation and software mapping','原生可编辑公式及代码变量映射；数字不因符号改变。'),
'18': ('已验证','正文以分析方法和目的叙述；探索性及日期记录留在补充与真实运行记录。','主稿 Methods；补充Notation and software mapping；执行记录','未伪造预注册或新增统计执行。'),
'19': ('已验证','明确D(t,k)为年度率差，ΔD为期末减期初，并明确主期间1990至2023。','主稿 Annual contrasts and the primary period change；公式1至6','同一亚型、期间、人口来源及单位；正分量提高有符号差。'),
'22': ('已验证','正文用机构和产品描述UN数据，确切产品文件与SHA256留在参考文献和来源清单。','主稿 Population sources；补充来源说明；输入清单','历史1950至2023产品中仅用1990至2023；UN只替换参照份额。完整Galaxy压缩源本轮未重新核验。'),
'25': ('已验证','方法分开结构敏感性与后验不确定性，列清已做敏感性及缺失信息，GATHER明确部分满足。','主稿 Structural sensitivity and posterior uncertainty；补充Joint uncertainty与GATHER','未将边际UI、模拟或国家bootstrap替代联合传播。'),
'26': ('已验证','使用原分析环境证据填写Python3.12.14、R4.6.1及NumPy/pandas/Matplotlib版本，并将文档环境另记。','主稿 Computational reproducibility；补充Notation and software mapping','这些库是外部依赖；未把当前文档机器版本冒充原科研环境。'),
'31': ('已验证','保留真实AI用途和已核实的早期模型标识，补记本轮日期；未知服务版本不猜测。','主稿 Methods与Acknowledgements；附信；AI证据JSON','作者最终核查和批准仍待作者确认，未伪称只润色。'),
'36': ('已验证','结果按覆盖、总体率差、三分量、亚型、UN和基准差异组织；复现细节移补充。','主稿 Results；原3表5图','GBD和UN数值分别标明；跨零仅称点估计符号改变。'),
'39': ('已验证','使用408行全年度亚型汇总与6行最近端点表回应抵消；补充明确2022两亚型同负、2023正负抵消。','补充Annual subtype differences；表S10；年度CSV','各年同病因/同分母/同标准/同基准加和通过；未把年份变化称显著相关。'),
'48': ('已验证','讨论直接说明国家份额对汇总比较的贡献及监测用途；经典方法归属集中在Methods。','主稿Discussion；补充S9/S11','不声称新算法、首次发现或纠正前人；临床资料只用于定义/背景。'),
'51': ('部分回应 受依赖限制','已完成访问审计并准确保留联合不确定性缺口；本轮未识别合格的联合draws和提供方配对证据。','主稿不确定性段；补充Joint uncertainty及GATHER；访问审计JSON','未生成后验UI，不能标统计解决。后续须取得真实样本和跨单元配对证据，重构摘要并逐draw重算及独立QC。'),
'76': ('已验证','主图5段可见Alt text及图片描述元数据保留；补充7图也有描述，逐图核对当前面板与主要方向。','主稿Figures1至5；补充FiguresS1至S7；Word图像元数据','符合保留现行可访问性说明的方案；不删除仅因它属于替代文本。')}


def responses():
    plan=json.loads((ROOT/'PCOS_Review_2026-10-30/20261030_review_plan_content.json').read_text())
    rows=[]
    for item in plan['comments']:
        cid=item['id'];status,action,evidence,reason=RESPONSE_ACTIONS[cid]
        rows.append({'ordinal':item['ordinal'],'comment_id':cid,'verbatim_comment':item['original'],'v5_status':item['v5_status'],'action':action,'reason':reason,'evidence':evidence,'output_path':str(TEX['main'].relative_to(ROOT))+';'+str(TEX['supplement'].relative_to(ROOT)),'verification':'content checks and page inspection required before final delivery','final_status':status})
    assert [r['comment_id'] for r in rows]==EXPECTED_IDS
    with (NEW/'docs'/(P+'v6_comment_response.csv')).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    save(NEW/'docs'/(P+'v6_comment_response.json'),rows)
    literature=[
        {'study':'Liu et al. 2024','doi':'10.1093/humrep/dead241','release_period':'GBD2019 1990-2019','target':'Geographical prevalence/YLD and trends','access':'Original Methods text via publisher search extraction and PubMed abstract; prior plan/full-text audit retained','relation':'Different target from explicit China-minus-18-country pool'},
        {'study':'Liang et al. 2025','doi':'10.3389/fpubh.2025.1693486','release_period':'GBD2021 1990-2021','target':'Burden/trends; Das Gupta population growth, ageing, epidemiological count components','access':'Publisher full-text Methods','relation':'Count components differ from country shares within fixed age standard'},
        {'study':'Liang 2026','doi':'10.1016/j.rbmo.2025.105239','release_period':'GBD2021 1990-2021; online2025 volume2026','target':'China/worldwide PCOS and related infertility trends, Joinpoint and ARIMA','access':'PubMed abstract only; full comparator construction unverified','relation':'No inference about an undocumented pooling formula'},
        {'study':'Yang et al. 2022','doi':'10.1016/j.lanwpc.2022.100494','release_period':'Community surveys2010/2020','target':'Rotterdam PCOS among married Han women20-49; conditional infertility characteristics','access':'PMC full text; detailed infertility questionnaire threshold not independently established','relation':'Definition context; no clinical rate validation'},
        {'study':'Teede et al. 2023','doi':'10.1210/clinem/dgad463','release_period':'2023 clinical guideline','target':'PCOS assessment and management','access':'PMC full text','relation':'Clinical diagnostic/care context; not a GBD modeled impairment-rate dataset'}]
    with (NEW/'docs'/(P+'v6_literature_estimands.csv')).open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(literature[0]));w.writeheader();w.writerows(literature)
    save(NEW/'docs'/(P+'v6_AI_evidence.json'),{'earlier_verified_evidence':json.loads((OLD/'docs/AI_VERSION_EVIDENCE.json').read_text()),'current_revision_date':'2026-10-30','current_tools':'OpenAI Codex; actual builder and QC commands retained','current_serving_model_identifier_independently_verified':False,'uses':['code development','analysis checks','drafting','technical editing','document production'],'author_final_approval':'pending; not assumed'})


def format_simple(doc):
    from docx.shared import Mm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    section=doc.sections[0];section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(25)
    for st in doc.styles:
        if st.type!=1:continue
        st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.line_spacing=1.15
        for b in st.element.xpath('.//w:pBdr'):b.getparent().remove(b)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control=True
        for r in p.runs:
            r.font.name='Times New Roman'
            fonts=r._element.get_or_add_rPr().rFonts
            if fonts is not None:fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    h=section.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if not h._p.xpath('.//w:fldSimple'):
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');h._p.append(field)


def documents():
    from docx import Document
    from docx.shared import Pt,Mm
    c=converter();c.build()
    # Cover letter is generated from the same editable TeX, retaining paragraph content.
    d=Document();format_simple(d)
    body=TEX['cover'].read_text().split(r'\begin{document}',1)[1].split(r'\end{document}')[0].replace(r'\sloppy','')
    for para in body.strip().split('\n\n'):
        content=re.sub(r'\n{2,}','\n',c.plain_tex(para.replace(r'\\','\n'))).strip()
        p=d.add_paragraph(content);p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.08
    for section in d.sections:
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(22)
    for p in d.paragraphs:
        for r in p.runs:r.font.size=Pt(10.5)
    d.save(WORD['cover'])
    # Separate editable main tables are copies of the verified native table objects.
    import copy
    main=Document(WORD['main']);d=Document();format_simple(d);d.add_paragraph('Tables for Human Reproduction v6','Title')
    mi=json.loads((NEW/'word/main_content_map.json').read_text())
    for i,(table,info) in enumerate(zip(main.tables,mi['tables'])):
        if i:d.add_page_break()
        d.add_paragraph('Table '+str(i+1)+' '+c.plain_tex(info['caption_tex']),'Heading 1')
        d._body._body.insert(len(d._body._body)-1,copy.deepcopy(table._tbl))
        d.add_paragraph(c.plain_tex(info['note_tex']))
    d.save(WORD['tables'])
    d=Document();format_simple(d)
    d.add_paragraph('PCOS manuscript responses to all 17 comments','Title')
    d.add_paragraph('Independent Human Reproduction v6 revision 30 October 2026','Subtitle')
    d.add_paragraph('逐条回应保留原始17条批注。16条已作修订或依据性回应；批注51仅完成访问审计及限制表述，联合不确定性传播仍受合格draws缺失限制。本文为作者审阅文件，未投稿。')
    rows=json.loads((NEW/'docs'/(P+'v6_comment_response.json')).read_text())
    for r in rows:
        d.add_paragraph('Comment '+r['comment_id'],'Heading 1')
        d.add_paragraph('原始批注','Heading 2');d.add_paragraph(r['verbatim_comment'])
        d.add_paragraph('处理结果 '+r['final_status'],'Heading 2');d.add_paragraph(r['action'])
        d.add_paragraph('依据与范围','Heading 2');d.add_paragraph(r['reason'])
        d.add_paragraph('核查位置','Heading 2');d.add_paragraph(r['evidence'])
    format_simple(d)
    for p in d.paragraphs:
        for r in p.runs:r.font.size=Pt(10.5 if p.style.name not in ['Title','Heading 1','Heading 2'] else 16 if p.style.name=='Title' else 12 if p.style.name=='Heading 1' else 11)
    d.save(WORD['response'])
    responsefonts()


def coverlayout():
    from docx import Document
    from docx.shared import Mm,Pt
    d=Document(WORD['cover'])
    for p in d.paragraphs:
        p.text=re.sub(r'\n{2,}','\n',p.text).strip()
        p.paragraph_format.line_spacing=1.08;p.paragraph_format.space_after=Pt(7)
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5)
    for section in d.sections:
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(22)
    d.save(WORD['cover'])
    verify()


def responsefonts():
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document(WORD['response'])
    for st in d.styles:
        if st.type!=1:continue
        st.font.name='Noto Sans CJK SC'
        pr=st.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Noto Sans CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    for p in d.paragraphs:
        pf=p.paragraph_format
        if p.style.name=='Heading 1':pf.space_before=Pt(10);pf.space_after=Pt(5)
        elif p.style.name=='Heading 2':pf.space_before=Pt(6);pf.space_after=Pt(3)
        elif p.style.name not in ['Title','Subtitle']:pf.space_after=Pt(4)
        for r in p.runs:
            r.font.name='Noto Sans CJK SC';pr=r._element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
            for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Noto Sans CJK SC')
            for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
            lang=OxmlElement('w:lang');lang.set(qn('w:val'),'zh-CN');lang.set(qn('w:eastAsia'),'zh-CN');pr.append(lang)
    d.save(WORD['response'])


def compile_tex():
    for kind,path in TEX.items():
        result=subprocess.run(['/opt/homebrew/bin/tectonic','--keep-logs','--outdir',str(NEW/'output'),str(path)],cwd=NEW/'manuscript',capture_output=True,text=True)
        write(PRIVATE/('compile_'+kind+'.log'),result.stdout+'\n'+result.stderr)
        assert result.returncode==0,(kind,result.stderr)
        # Tectonic uses the complete dated source stem as the PDF stem.
        assert PDF[kind].exists()


def refresh():
    text=TEX['main'].read_text()
    a=text.index(r'\noindent\textbf{STUDY QUESTION:}')
    b=text.index(r'\noindent\textbf{FUNDING:}',a)
    write(TEX['main'],text[:a]+ABSTRACT+text[b:])
    diff=''.join(difflib.unified_diff((OLD/'manuscript/main.tex').read_text().splitlines(True),TEX['main'].read_text().splitlines(True),fromfile='v5/main.tex',tofile=TEX['main'].name))
    write(NEW/'docs'/(P+'v5_to_v6_main.diff'),diff)
    documents();compile_tex();verify()
    shutil.copy2(NEW/'docs'/(P+'v6_content_verification.json'),NEW/'docs'/(P+'v6_retry_verification.json'))


def verify():
    from docx import Document
    from lxml import etree
    c=converter();c.verify()
    baseline=json.loads((NEW/'docs'/(P+'v5_preservation_baseline.json')).read_text())
    assert all(sha(ROOT/p)==h for p,h in baseline.items())
    maps={kind:json.loads((NEW/'word'/(kind+'_content_map.json')).read_text()) for kind in ['main','supplement']}
    for kind in maps:
        before=json.loads((OLD/'word'/(kind+'_content_map.json')).read_text());after=maps[kind]
        assert [f['sha256'] for f in before['figures']]==[f['sha256'] for f in after['figures']]
        for a,b in zip(before['tables'][:3] if kind=='main' else before['tables'][1:9],after['tables'][:3] if kind=='main' else after['tables'][1:9]):
            assert a['rows']==b['rows'],(kind,a['number'])
        assert before['xx_count']==after['xx_count']
    main=TEX['main'].read_text();supp=TEX['supplement'].read_text()
    assert r'C_{catk}' not in main and r'A_{Ctk}' not in main and r'\sum_{c\in R}' not in main
    assert all(s in main for s in [r'N_{c,a,t,k}',r'\mathrm{CHN}',r'\Delta D_k(t_0,t_1)=D_{t_1,k}-D_{t_0,k}'])
    assert 'in this journal' not in main.lower()
    for v in ['271.51','208.62','88.65','151.54','86.76','2.16','2.26','48.34']:
        assert v in main,v
    assert len(maps['main']['numbered_equations'])==7
    assert len(maps['main']['tables'])==3 and len(maps['supplement']['tables'])==12
    assert main.count(r'\textbf{Alt text:}')==5
    for kind,oldname in [('main','main'),('supplement','supplement'),('cover','cover_letter')]:
        assert TEX[kind].read_text().count('XX')==(OLD/'manuscript'/(oldname+'.tex')).read_text().count('XX')
    # Same counting method for both abstracts: Pandoc plain, whitespace words.
    counts={}
    for version,text in [('v5',(OLD/'manuscript/main.tex').read_text()),('v6',main)]:
        abstract=text.split(r'\begin{abstract}')[1].split(r'\end{abstract}')[0]
        plain=c.plain_tex(abstract)
        heads=re.findall(r'\\textbf\{([^}]+):\}',abstract)
        without=re.sub(r'\\textbf\{[^}]+:\}','',abstract)
        counts[version]={'words_without_headings':len(c.plain_tex(without).split()),'words_with_headings':len(plain.split()),'headings':heads}
        assert len(heads)==11
    assert counts['v6']['words_without_headings']<counts['v5']['words_without_headings']
    assert 350<=counts['v6']['words_without_headings']<=450,counts
    save(NEW/'docs'/(P+'v6_abstract_count.json'),{'method':'Pandoc LaTeX-to-plain followed by whitespace split; same method both versions','internal_target_without_headings':[350,450],'journal_hard_limit_asserted':False,'counts':counts})
    rows=json.loads((NEW/'docs'/(P+'v6_comment_response.json')).read_text())
    plan=json.loads((ROOT/'PCOS_Review_2026-10-30/20261030_review_plan_content.json').read_text())
    assert [r['comment_id'] for r in rows]==EXPECTED_IDS
    assert [r['verbatim_comment'] for r in rows]==[r['original'] for r in plan['comments']]
    response=Document(WORD['response']);all_response='\n'.join(p.text for p in response.paragraphs)
    assert all(r['verbatim_comment'] in all_response for r in rows)
    assert '部分回应 受依赖限制' in all_response
    # Cover content must match the source's paragraph stream, apart from line breaks.
    cover=Document(WORD['cover']);covertext=''.join(p.text for p in cover.paragraphs)
    source=TEX['cover'].read_text().split(r'\begin{document}',1)[1].split(r'\end{document}')[0].replace(r'\sloppy','')
    assert re.sub(r'\s+','',covertext)==re.sub(r'\s+','',c.plain_tex(source.replace(r'\\','\n')))
    assert covertext.count('XX')==5
    from pypdf import PdfReader
    pdfchecks={}
    for kind in TEX:
        p=PDF[kind];reader=PdfReader(p);text='\n'.join(page.extract_text() or '' for page in reader.pages)
        assert 'XX' in text
        assert '\ufffd' not in text
        assert not re.search(r'undefined (?:reference|citation)',text,re.I)
        log=(PRIVATE/('compile_'+kind+'.log')).read_text()
        assert 'undefined reference' not in log.lower() and 'undefined citation' not in log.lower()
        pdfchecks[kind]={'pages':len(reader.pages),'sha256':sha(p),'compile_exit_code':0,'visual_inspection':'pending'}
    save(NEW/'docs'/(P+'v6_content_verification.json'),{'verified_at':datetime.now(ZoneInfo('Europe/Paris')).isoformat(timespec='seconds'),'old_v5_files_preserved':len(baseline),'unchanged_core_tables':3,'unchanged_existing_supplement_numerical_tables':8,'new_supplement_tables':['S10 recent endpoints','S11 outcome definitions'],'main_numbered_equations':7,'all_17_verbatim_comments_preserved':True,'comments_with_completed_edit_or_reasoned_response':16,'comment_51':'partial response; qualified joint draws missing','core_point_estimates_changed':False,'new_scientific_model_fitted':False,'posterior_intervals_added':False,'un_source_role_preserved':True,'original_XX_counts':{k:TEX[k].read_text().count('XX') for k in TEX},'pdfs':pdfchecks,'visual_inspection':'pending','submitted':False})
    print(json.dumps({'content_checks':'passed','abstract_words':counts,'pdfs':pdfchecks},ensure_ascii=False))


def package():
    qa=json.loads((NEW/'docs'/(P+'v6_visual_QA.json')).read_text())
    assert qa['all_pages_visually_inspected'] and qa['unresolved_layout_defects']==0
    # Restore helper names for the converter's content verification on a retry.
    for dated in list((NEW/'word').glob(P+'v6_*.json')):
        dated.rename(dated.with_name(dated.name[len(P+'v6_'):]))
    verify()
    final=json.loads((NEW/'docs'/(P+'v6_content_verification.json')).read_text())
    final['visual_inspection']='passed; every Word-rendered and TeX PDF page inspected';final['total_pages_inspected']=qa['total_pages_inspected']
    final['verification_scope']='Content/math/values/placeholders/pages and preservation; author final scientific approval remains pending'
    save(NEW/'docs'/(P+'v6_FINAL_VERIFICATION.json'),final)
    # Rename conversion-support records to dated names before public delivery.
    for p in list((NEW/'word').iterdir()):
        if p.is_file() and not p.name.startswith(P):p.rename(p.with_name(P+'v6_'+p.name))
    write(NEW/(P+'README_ZH.txt'),'''Human Reproduction 独立v6 2026-10-30

按全部17条批注修订主稿、补充稿、附信及逐条回应；保留v5与原XX。16条完成编辑或依据性回应；批注51仅部分回应，缺少合格联合draws，未生成派生后验区间。作者最终核查/批准和既定XX仍待投稿阶段由作者处理，本轮暂不投稿。

word目录提供完整可编辑主稿、补充稿、附信、逐条回应和独立主表；output目录提供TeX主稿/补充/附信PDF及Word生成的逐条回应PDF；manuscript目录含对应日期前缀TeX与原数值图表。Word与TeX的分页可以不同，内容和数值已核对。docs目录含17行原批注回应、来源/方法/文献/AI/差异/数值与逐页视觉核验。results含408行全年亚型率差汇总；没有国家年龄GBD源单元。

主要变化：病例N/中国CHN与年度D/期间ΔD统一；摘要与贡献表达精简；调查五年/临床12个月/全女性模型分母说明；G20固定19国/18参照及真实环境；补充S10年度抵消与S11定义对照。原3主表、8补充数值表和12图内容保留。数字来自既有科学快照和已验证前置核查，未重拟合GBD或升级权重。官方全精度标准权重向量来源仍未独立确认。

主分析代码与结果固定科学快照1abf337d9f23cc69292e03cd581b6a1557e6effe。构建代码scripts/build_hr_v6_20261030.py复用已有Word转换器；本次真实命令/环境/输入与输出哈希见共享CODEX_RUN_LOG、构建来源及最终核验。完整包只含允许公开的文稿、代码、汇总与记录，受限GBD原件和医院资料不在内。
''')
    files=[p for p in NEW.rglob('*') if p.is_file() and p.suffix not in ['.zip','.log'] and p.name!=P+'v6_DELIVERY_MANIFEST.json']
    files += [Path(__file__).resolve(),ROOT/'scripts/convert_hr_v4_word_20261029.py',ROOT/'scripts/render_hr_v6_20261030.py']
    manifest={'files':[{'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':sha(p),'git_blob':subprocess.check_output(['git','hash-object',str(p)],cwd=ROOT,text=True).strip()} for p in sorted(files)],'restricted_source_data_included':False,'submitted':False}
    save(NEW/'docs'/(P+'v6_DELIVERY_MANIFEST.json'),manifest)
    files.append(NEW/'docs'/(P+'v6_DELIVERY_MANIFEST.json'))
    archive=NEW/'output'/(BASE+'Complete.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(files):z.write(p,p.relative_to(ROOT))
    print(json.dumps({'archive':str(archive.relative_to(ROOT)),'members':len(files),'sha256':sha(archive)}))


def verify_package():
    manifest_path=NEW/'docs'/(P+'v6_DELIVERY_MANIFEST.json')
    manifest=json.loads(manifest_path.read_text())
    archive=NEW/'output'/(BASE+'Complete.zip')
    with zipfile.ZipFile(archive) as z:
        expected={r['path'] for r in manifest['files']}|{str(manifest_path.relative_to(ROOT))}
        assert set(z.namelist())==expected
        for r in manifest['files']:
            assert sha(ROOT/r['path'])==r['sha256']
            assert hashlib.sha256(z.read(r['path'])).hexdigest()==r['sha256']
    baseline=json.loads((NEW/'docs'/(P+'v5_preservation_baseline.json')).read_text())
    assert all(sha(ROOT/p)==h for p,h in baseline.items())
    assert all(p.name.startswith(P) for p in NEW.rglob('*') if p.is_file() and p.suffix!='.log')
    from pypdf import PdfReader
    response='\n'.join(p.extract_text() or '' for p in PdfReader(NEW/'output'/(BASE+'CommentResponse.pdf')).pages)
    assert '原始批注' in response and '联合' in response and '\ufffd' not in response
    report={'package_check':'passed','manifest_files':len(manifest['files']),'zip_members':len(expected),'v5_files_preserved':len(baseline),'archive_sha256':sha(archive)}
    save(NEW/'docs'/(P+'v6_PACKAGE_CHECK.json'),report)
    print(json.dumps(report))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('action',choices=['source','documents','compile','verify','build','refresh','coverlayout','responsefonts','package','verify_package']);args=parser.parse_args()
    PRIVATE.mkdir(exist_ok=True,parents=True)
    if args.action=='build':sources();documents();compile_tex();verify()
    else:{'source':sources,'documents':documents,'compile':compile_tex,'verify':verify,'refresh':refresh,'coverlayout':coverlayout,'responsefonts':responsefonts,'package':package,'verify_package':verify_package}[args.action]()
