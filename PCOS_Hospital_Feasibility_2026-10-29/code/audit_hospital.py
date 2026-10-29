#!/usr/bin/env python3
"""Read hospital files in place; export aggregate feasibility statistics only.

Usage: python audit_hospital.py --source-dir <private folder> --output <summary.json>
No patient names, identifiers, birth dates, contact details or row extracts exported.
"""
from pathlib import Path
from collections import Counter
import argparse
import datetime as dt
import hashlib
import json
import re
import numpy as np
import pandas as pd
import openpyxl
import xlrd

SENSITIVE=re.compile('姓名|门诊号|身份证|电话|地址|生日|出生|患者编码|标本号|条码|审核者|开立医生|床号')
PCOS=re.compile(r'PCOS|多囊',re.I)

def blank(v):
    return v is None or (isinstance(v,str) and not v.strip()) or (isinstance(v,float) and np.isnan(v))

def text(v):
    if blank(v):return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).strip()

def numeric(v):
    if isinstance(v,(int,float)):return float(v)
    s=text(v)
    if re.fullmatch(r'[+-]?\d+(\.\d+)?',s):return float(s)
    return np.nan

def day(v):
    if isinstance(v,(dt.datetime,dt.date)):return v.date() if isinstance(v,dt.datetime) else v
    s=text(v)
    if re.fullmatch(r'\d{8}',s):
        z=pd.to_datetime(s,format='%Y%m%d',errors='coerce')
        return z.date() if not pd.isna(z) else None
    if re.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}',s):
        return pd.to_datetime(s,errors='coerce').date()
    return None

def safe_vocab(values):
    c=Counter(text(x) for x in values if not blank(x))
    out={k:n for k,n in c.items() if n>=5 and len(k)<=16 and not re.search(r'\d{6}|\n|\r',k)}
    return {'common_coded_values_n_ge_5':out,'other_nonempty_records':sum(c.values())-sum(out.values()),'distinct_nonempty_values':len(c)}

def xlsx_read(p):
    w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    frames=[];sheets=[]
    for sh in w:
        rows=list(sh.iter_rows(values_only=True))
        rows=[r for r in rows if any(not blank(x) for x in r)]
        sheets.append({'sheet':sh.title,'nonempty_rows_including_header':len(rows),'declared_rows':sh.max_row,'declared_columns':sh.max_column})
        if len(rows)<2:continue
        hdr=[text(v) or f'Unnamed_{i+1}' for i,v in enumerate(rows[0])]
        seen=Counter();cols=[]
        for h in hdr:
            seen[h]+=1;cols.append(h if seen[h]==1 else h+f'__repeat{seen[h]}')
        f=pd.DataFrame(rows[1:],columns=cols);frames.append(f)
    w.close()
    assert len(frames)==1
    return frames[0],sheets

def audit_clinical(p):
    d,sheets=xlsx_read(p);n=len(d)
    nonmissing={c:sum(not blank(v) for v in d[c]) for c in d}
    names=d['女姓名'].map(text)
    namecounts=names[names!=''].value_counts()
    result={'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':sheets,'records':n,'columns':list(d.columns),'nonmissing_by_column':nonmissing,'sensitive_columns':[c for c in d if SENSITIVE.search(c)],'distinct_female_name_tokens':len(namecounts),'female_name_tokens_with_multiple_rows':int((namecounts>1).sum()),'rows_beyond_first_by_female_name':int((namecounts-1).clip(lower=0).sum()),'name_deduplication_is_not_verified_patient_identity':True}
    if '女门诊号' in d:
        ids=d['女门诊号'].map(text);ids=ids[ids!=''];vc=ids.value_counts()
        result['nonmissing_female_outpatient_ID_records']=len(ids)
        result['distinct_female_outpatient_IDs']=len(vc)
        result['ID_tokens_with_multiple_rows']=int((vc>1).sum())
        result['rows_beyond_first_by_ID']=int((vc-1).clip(lower=0).sum())
        result['unique_ID_and_cycle_date_pairs']=int(d.assign(_id=d['女门诊号'].map(text),_date=d['周期开始日期'].map(text)).drop_duplicates(['_id','_date']).shape[0])
        result['same_ID_multiple_name_tokens']=sum(d.assign(_id=d['女门诊号'].map(text),_name=names).groupby('_id')['_name'].nunique()>1)
    diag=d['女方病历诊断'].map(text)
    result['PCOS_text_flag_records']=int(diag.str.contains(PCOS).sum())
    result['primary_infertility_text_flag_records']=int(diag.str.contains('原发|原发性').sum())
    result['secondary_infertility_text_flag_records']=int(diag.str.contains('继发|继发性').sum())
    result['outcome_fields_present']=[c for c in d if re.search('妊娠|流产|移植|随访|活产|分娩|新生儿|OHSS',c)]
    for c in result['outcome_fields_present']:result[c+'_coding']=safe_vocab(d[c]) if '日期' not in c else {'nonmissing':nonmissing[c]}
    result['numeric_summaries']={}
    for c in ['女年龄','女方BMI','不孕年限','总AFC数','左侧AFC','右侧AFC','婚育史孕次','婚育史产次']:
        if c not in d:continue
        v=d[c].map(numeric).dropna()
        result['numeric_summaries'][c]={'numeric_n':len(v),'unparsed_nonempty':nonmissing[c]-len(v),'min':float(v.min()) if len(v) else None,'median':float(v.median()) if len(v) else None,'max':float(v.max()) if len(v) else None}
    result['duplicate_header_consistency']={}
    G=d['婚育史孕次'].map(numeric);P=d['婚育史产次'].map(numeric)
    result['history_field_checks']={
        'parity_gt_gravidity_if_headers_interpreted_literally':int((P>G).sum()),
        'primary_label_with_positive_gravidity':int((diag.str.contains('原发')&(G>0)).sum()),
        'secondary_label_with_zero_gravidity':int((diag.str.contains('继发')&(G==0)).sum()),
        'AFC_numeric_above_100':{c:int((d[c].map(numeric)>100).sum()) for c in ['左侧AFC','右侧AFC']},
        'interpretation':'Potential label/encoding problems requiring source clarification; no automatic recoding.'}
    if '是否妊娠' in d:
        tmp=d.assign(_transfer=d['是否移植'].map(text),_observed=d['是否妊娠'].map(lambda x:not blank(x)))
        result['transfer_and_outcome']=[{'transfer':a,'outcome_recorded':bool(b),'records':int(n)} for (a,b),n in tmp.groupby(['_transfer','_observed']).size().items()]
        outcomes=d[d['是否妊娠'].map(lambda x:not blank(x))]
        result['outcome_distinct_nonempty_female_IDs']=int(outcomes['女门诊号'].map(text).replace('',None).nunique())
        result['outcome_distinct_female_name_tokens']=int(outcomes['女姓名'].map(text).nunique())
    for c in ['基础激素FSH','基础激素LH','基础激素E2','基础激素P','基础激素PRL','基础激素T']:
        cc=c+'__repeat2'
        if cc not in d:continue
        both=d[c].map(lambda x:not blank(x)) & d[cc].map(lambda x:not blank(x))
        result['duplicate_header_consistency'][c]={'both_nonempty':int(both.sum()),'equal_text':int((d.loc[both,c].map(text)==d.loc[both,cc].map(text)).sum())}
    if '周期开始日期' in d:
        dates=[day(v) for v in d['周期开始日期']];valid=[v for v in dates if v is not None and not pd.isna(v)]
        result['cycle_start_year_range']=[min(valid).year,max(valid).year] if valid else None
        result['cycle_start_nonmissing']=nonmissing['周期开始日期']
    d['_name']=names
    return result,d

def audit_lab(p):
    b=xlrd.open_workbook(p,on_demand=True);assert b.nsheets==1;s=b.sheet_by_index(0)
    hdr=[text(v) for v in s.row_values(3)]
    rows=[]
    for i in range(4,s.nrows):
        row=s.row_values(i)
        if any(not blank(v) for v in row):rows.append(row)
    d=pd.DataFrame(rows,columns=hdr)
    assert '患者编码' in d and '项目名称' in d
    ids=d['患者编码'].map(text);names=d['患者姓名'].map(text)
    numeric_result=d['结果'].map(numeric)
    dates=[]
    for i,v in enumerate(d['检验日期']):
        if isinstance(v,(int,float)) and 20000<float(v)<70000:
            dates.append(xlrd.xldate_as_datetime(v,b.datemode).date())
        else:dates.append(day(v))
    valid=[v for v in dates if v is not None and not pd.isna(v)]
    both=ids!=''
    result={'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'records':len(d),'columns':hdr,'sensitive_columns':[c for c in hdr if SENSITIVE.search(c)],'nonmissing_by_column':{c:sum(not blank(v) for v in d[c]) for c in d},'distinct_patient_codes':int(ids[ids!=''].nunique()),'distinct_female_patient_codes':int(ids[(d['患者性别'].map(text).isin(['女','女性','Female','F']))&(ids!='')].nunique()),'sex_codes':safe_vocab(d['患者性别']),'distinct_name_tokens':int(names[names!=''].nunique()),'numeric_result_records':int(numeric_result.notna().sum()),'test_date_year_month_range':[min(valid).strftime('%Y-%m'),max(valid).strftime('%Y-%m')] if valid else None,'PCOS_text_flag_lab_rows':int(d['诊断'].map(text).str.contains(PCOS).sum()),'PCOS_text_flag_patient_codes':int(ids[d['诊断'].map(text).str.contains(PCOS)&(ids!='')].nunique())}
    result['named_test_rows']=int(d['项目名称'].map(text).ne('').sum())
    result['missing_patient_code_rows']=int(ids.eq('').sum())
    result['tests']=[]
    for (test,unit),g in d.groupby(['项目名称','单位'],dropna=False):
        codes=g['患者编码'].map(text)
        result['tests'].append({'test':text(test),'unit':text(unit),'records':len(g),'distinct_patient_codes':int(codes[codes!=''].nunique()),'numeric_results':int(g['结果'].map(numeric).notna().sum())})
    result['tests']=sorted(result['tests'],key=lambda x:-x['records'])
    d['_name']=names;d['_code']=ids
    d['_fingerprint']=d.apply(lambda r:hashlib.sha256('\x1f'.join(text(r.get(c)) for c in ['患者编码','标本号','检验日期','项目编码','项目名称','结果','单位']).encode()).hexdigest(),axis=1)
    b.release_resources()
    return result,d

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--source-dir',required=True);ap.add_argument('--output',required=True);args=ap.parse_args()
    source=Path(args.source_dir);reports=[];clinical={};labs={}
    for p in sorted(source.glob('*.xlsx')):
        r,d=audit_clinical(p);reports.append(r);clinical[p.name]=d
    for p in sorted(source.glob('*.xls')):
        r,d=audit_lab(p);reports.append(r);labs[p.name]=d
    overlaps=[]
    if len(labs)==2:
        (pa,aa),(pb,bb)=list(labs.items());sa=set(aa['_fingerprint']);sb=set(bb['_fingerprint'])
        overlaps.append({'files':[pa,pb],'exact_lab_row_fingerprint_overlap':len(sa&sb),'unique_lab_row_fingerprints_union':len(sa|sb),'within_file_repeated_fingerprints':[len(aa)-len(sa),len(bb)-len(sb)],'shared_patient_code_tokens':len(set(aa['_code'])-{''}&(set(bb['_code'])-{''}))})
    linkage=[]
    for cn,cc in clinical.items():
        for ln,ll in labs.items():
            names=set(cc['_name'])-{''};labnames=set(ll['_name'])-{''}
            record={'clinical_file':cn,'lab_file':ln,'overlapping_name_tokens_only':len(names&labnames),'not_a_validated_patient_linkage':True}
            if '女门诊号' in cc:
                record['exact_outpatient_ID_vs_lab_code_token_overlap']=len((set(cc['女门诊号'].map(text))-{''})&(set(ll['_code'])-{''}))
                idnames=ll.groupby('_code')['_name'].agg(set).to_dict()
                shared=cc[cc['女门诊号'].map(text).isin(set(ll['_code'])-{''})]
                record['female_rows_shared_ID_and_name_concordant']=int(sum(text(row['女姓名']) in idnames.get(text(row['女门诊号']),set()) for _,row in shared.iterrows()))
                record['female_shared_ID_records']=len(shared)
                record['male_ID_vs_lab_code_token_overlap']=len((set(cc['男门诊号'].map(text))-{''})&(set(ll['_code'])-{''}))
                sharedmale=cc[cc['男门诊号'].map(text).isin(set(ll['_code'])-{''})]
                record['male_rows_shared_ID_and_name_concordant']=int(sum(text(row['男姓名']) in idnames.get(text(row['男门诊号']),set()) for _,row in sharedmale.iterrows()))
                record['male_shared_ID_records']=len(sharedmale)
            linkage.append(record)
    output={'date':'2026-10-29','scope':'Read-only feasibility audit; no outcome association tests or clinical model fit. No patient-level data exported.','files':reports,'lab_overlap':overlaps,'linkage_diagnostics':linkage,'limits':['Column labels and diagnosis text do not establish Rotterdam PCOS eligibility or outcome definitions.','Name-token duplicates are not verified patient identities.','Exact ID token comparisons do not prove or exclude linkage across incompatible hospital systems.','No individual-level source data are suitable for a public repository.']}
    Path(args.output).write_text(json.dumps(output,indent=2,ensure_ascii=False,default=lambda x:int(x) if isinstance(x,np.integer) else str(x)))
    for x in reports:
        print(x['file'],'records',x['records'],'PCOS rows',x.get('PCOS_text_flag_records',x.get('PCOS_text_flag_lab_rows')))
    print('Aggregate audit written; no patient rows exported.')

if __name__=='__main__':main()
